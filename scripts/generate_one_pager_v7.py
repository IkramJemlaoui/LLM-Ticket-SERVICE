"""Generate the six-section AegisDesk agentic application one-pager."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "AegisDesk_AI_One_Pager_v7.pptx"

NAVY = RGBColor(10, 47, 90)
NAVY_2 = RGBColor(19, 57, 92)
CYAN = RGBColor(27, 181, 213)
BLUE = RGBColor(37, 91, 173)
VIOLET = RGBColor(113, 55, 166)
TEAL = RGBColor(7, 143, 145)
ORANGE = RGBColor(231, 138, 18)
PURPLE = RGBColor(102, 69, 170)
GREEN = RGBColor(75, 148, 30)
TEXT = RGBColor(29, 55, 82)
MUTED = RGBColor(78, 101, 125)
LINE = RGBColor(184, 207, 226)
PALE = RGBColor(246, 250, 253)
WHITE = RGBColor(255, 255, 255)


def textbox(slide, x, y, w, h, text, size=8, color=TEXT, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.025)
    frame.margin_top = frame.margin_bottom = Inches(0.015)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def rounded_box(slide, x, y, w, h, fill, line=LINE, radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def card(slide, x, y, w, h, number, title, accent):
    rounded_box(slide, x, y, w, h, WHITE, LINE)
    marker = rounded_box(slide, x + 0.11, y + 0.1, 0.31, 0.31, accent, accent, MSO_SHAPE.OVAL)
    marker.text_frame.clear()
    marker.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = marker.text_frame.paragraphs[0]
    p.text = str(number)
    p.font.name = "Aptos Display"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    textbox(slide, x + 0.5, y + 0.1, w - 0.62, 0.3, title.upper(), 8.6, accent, True, "Aptos Display")


def icon_row(slide, x, y, w, icon, question, answer, accent, answer_size=6.6):
    textbox(slide, x, y + 0.02, 0.28, 0.3, icon, 13, accent, False, "Segoe UI Emoji", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    textbox(slide, x + 0.34, y, 1.05, 0.35, question, 5.9, accent, True, "Aptos", PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    box = rounded_box(slide, x + 1.43, y, w - 1.43, 0.36, PALE, RGBColor(222, 233, 241))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = Inches(0.07)
    box.text_frame.margin_top = box.text_frame.margin_bottom = Inches(0.02)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = box.text_frame.paragraphs[0]
    p.text = answer
    p.font.name = "Aptos"
    p.font.size = Pt(answer_size)
    p.font.color.rgb = TEXT


def agent_node(slide, cx, cy, label, icon, accent):
    node = rounded_box(slide, cx - 0.22, cy - 0.22, 0.44, 0.44, WHITE, accent, MSO_SHAPE.OVAL)
    node.line.width = Pt(1.2)
    node.text_frame.clear()
    node.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = node.text_frame.paragraphs[0]
    p.text = icon
    p.font.name = "Segoe UI Emoji"
    p.font.size = Pt(11)
    p.alignment = PP_ALIGN.CENTER
    textbox(slide, cx - 0.42, cy + 0.23, 0.84, 0.19, label, 4.8, accent, True, "Aptos", PP_ALIGN.CENTER)


def connector(slide, x1, y1, x2, y2, color):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(1.0)
    return line


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(241, 248, 253)

    # Header and purpose strip.
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.72))
    header.fill.solid(); header.fill.fore_color.rgb = NAVY; header.line.fill.background()
    textbox(slide, 0.2, 0.11, 10.7, 0.32, "LLM & BUSINESS APPLICATIONS — AGENTIC APPLICATION DESIGN", 16, WHITE, True, "Aptos Display")
    textbox(slide, 0.2, 0.43, 7.0, 0.2, "AEGISDESK AI · SIX-SECTION ONE-PAGER", 8.1, CYAN, True)
    logo = rounded_box(slide, 11.73, 0.12, 1.36, 0.45, NAVY_2, CYAN)
    logo.text_frame.clear(); logo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = logo.text_frame.paragraphs[0]; p.text = "AEGISDESK AI"; p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos Display"; p.font.size = Pt(9); p.font.bold = True; p.font.color.rgb = WHITE

    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.79), Inches(12.89), Inches(0.34))
    strip.fill.solid(); strip.fill.fore_color.rgb = RGBColor(230, 243, 250); strip.line.color.rgb = LINE
    textbox(slide, 0.3, 0.84, 0.3, 0.22, "💡", 11, PURPLE, False, "Segoe UI Emoji", PP_ALIGN.CENTER)
    textbox(slide, 0.65, 0.84, 12.15, 0.22, "One intelligent front door that clarifies, routes and resolves cross-department requests with grounded AI and human control.", 7.3, TEXT)

    margin, gap = 0.2, 0.11
    width = (13.333 - 2 * margin - 2 * gap) / 3
    y_top, y_bottom, height = 1.22, 3.82, 2.47
    xs = [margin, margin + width + gap, margin + 2 * (width + gap)]

    # 1 — Application Overview.
    card(slide, xs[0], y_top, width, height, 1, "Application Overview", BLUE)
    rows = [
        ("🏷️", "Application name", "AegisDesk AI"),
        ("👥", "Target users", "Every employee; every department can request or resolve."),
        ("🎯", "Context & problem", "Incomplete requests, wrong routing and scattered company knowledge."),
        ("💎", "Value proposition", "Faster guidance, complete cases, correct ownership and reusable solutions."),
    ]
    for index, values in enumerate(rows):
        icon_row(slide, xs[0] + 0.14, y_top + 0.55 + index * 0.45, width - 0.28, *values, BLUE)

    # 2 — Challenges.
    card(slide, xs[1], y_top, width, height, 2, "Key Challenges Addressed", VIOLET)
    challenges = [
        ("⚠️", "Challenge 1", "Missing facts create clarification loops."),
        ("🔀", "Challenge 2", "Wrong routing delays the accountable team."),
        ("📚", "Challenge 3", "Answers are scattered across documents, cases and people."),
        ("🛡️", "Challenge 4", "Unsupported AI creates trust, privacy and safety risk."),
    ]
    for index, values in enumerate(challenges):
        icon_row(slide, xs[1] + 0.14, y_top + 0.55 + index * 0.45, width - 0.28, *values, VIOLET)

    # 3 — Capabilities.
    card(slide, xs[2], y_top, width, height, 3, "Core Capabilities (Services Delivered)", TEAL)
    capabilities = [
        ("📝", "Capability 1", "Clarify the request before submission."),
        ("🧭", "Capability 2", "Classify, prioritise and recommend the responsible team."),
        ("🔎", "Capability 3", "Find approved guidance and similar verified cases."),
        ("✅", "Capability 4", "Draft a cited answer for human review and controlled release."),
    ]
    for index, values in enumerate(capabilities):
        icon_row(slide, xs[2] + 0.14, y_top + 0.55 + index * 0.45, width - 0.28, *values, TEAL)

    # 4 — Differentiation.
    card(slide, xs[0], y_bottom, width, height, 4, "Differentiation", ORANGE)
    differentiation = [
        ("✨", "What is unique?", "One platform for requesters and resolver teams across the company."),
        ("🤖", "Agentic value", "Five bounded specialists create visible, testable handoffs."),
        ("🔌", "Integration", "Designed for company knowledge, case systems, SSO and collaboration tools."),
        ("📈", "Scale", "Permission-aware retrieval, monitoring and reusable verified solutions."),
    ]
    for index, values in enumerate(differentiation):
        icon_row(slide, xs[0] + 0.14, y_bottom + 0.55 + index * 0.45, width - 0.28, *values, ORANGE, 6.3)

    # 5 — Agentic approach, with a visual inspired by the reference template.
    card(slide, xs[1], y_bottom, width, height, 5, "Agentic Approach (At a Glance)", BLUE)
    qx = xs[1] + 0.15
    questions = [
        ("How many agents?", "Five specialists"),
        ("Main roles?", "Safety · Triage · Knowledge · Resolution · Quality"),
        ("Collaboration?", "Deterministic orchestrator"),
        ("Human-in-loop?", "Approve · edit · reassign · escalate"),
    ]
    for index, (question, answer) in enumerate(questions):
        yy = y_bottom + 0.57 + index * 0.39
        textbox(slide, qx, yy, 0.92, 0.2, question, 5.2, BLUE, True, valign=MSO_ANCHOR.MIDDLE)
        textbox(slide, qx + 0.92, yy, 0.8, 0.27, answer, 5.2, TEXT, False, valign=MSO_ANCHOR.MIDDLE)

    center_x, center_y = xs[1] + 3.05, y_bottom + 1.31
    node_points = [
        (xs[1] + 2.28, y_bottom + 0.78, "Safety", "🛡️"),
        (xs[1] + 3.02, y_bottom + 0.66, "Triage", "🧭"),
        (xs[1] + 3.71, y_bottom + 1.02, "Knowledge", "📚"),
        (xs[1] + 3.52, y_bottom + 1.78, "Resolution", "✍️"),
        (xs[1] + 2.56, y_bottom + 1.85, "Quality", "✅"),
    ]
    for nx, ny, _, _ in node_points:
        connector(slide, center_x, center_y, nx, ny, RGBColor(129, 166, 198))
    hub = rounded_box(slide, center_x - 0.3, center_y - 0.3, 0.6, 0.6, RGBColor(231, 243, 252), BLUE, MSO_SHAPE.OVAL)
    hub.line.width = Pt(1.3); hub.text_frame.clear(); hub.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = hub.text_frame.paragraphs[0]; p.text = "🤖"; p.font.name = "Segoe UI Emoji"; p.font.size = Pt(15); p.alignment = PP_ALIGN.CENTER
    textbox(slide, center_x - 0.43, center_y + 0.3, 0.86, 0.18, "ORCHESTRATOR", 4.6, BLUE, True, align=PP_ALIGN.CENTER)
    for nx, ny, label, icon in node_points:
        agent_node(slide, nx, ny, label, icon, BLUE)
    human = rounded_box(slide, xs[1] + 2.17, y_bottom + 2.15, 1.92, 0.23, RGBColor(244, 239, 253), PURPLE)
    human.text_frame.clear(); human.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = human.text_frame.paragraphs[0]; p.text = "👤 HUMAN APPROVAL · FINAL AUTHORITY"; p.font.name = "Aptos"; p.font.size = Pt(5.3); p.font.bold = True; p.font.color.rgb = PURPLE; p.alignment = PP_ALIGN.CENTER

    # 6 — Expected Business Impact.
    card(slide, xs[2], y_bottom, width, height, 6, "Expected Business Impact", PURPLE)
    impact = [
        ("🎯", "Business impact", "Fewer delays and transfers; faster, accountable cross-team service."),
        ("📊", "KPIs", "Pilot: ≥85% complete/routed · −40% draft time · 0 unapproved sends."),
        ("💰", "Value / ROI", "Time saved × eligible volume × loaded cost − operating cost."),
        ("⏱️", "Time to value", "Four-week pilot; scale only after measured quality, safety and value."),
    ]
    for index, values in enumerate(impact):
        icon_row(slide, xs[2] + 0.14, y_bottom + 0.55 + index * 0.45, width - 0.28, *values, PURPLE, 6.2)

    # Footer: project stack, disclosure and student information.
    footer = rounded_box(slide, 0.2, 6.4, 12.89, 0.77, WHITE, LINE)
    textbox(slide, 0.34, 6.49, 0.28, 0.28, "🧩", 12, TEAL, font="Segoe UI Emoji", align=PP_ALIGN.CENTER)
    textbox(slide, 0.67, 6.47, 7.95, 0.2, "PROJECT FOUNDATION", 6.2, TEAL, True)
    textbox(slide, 0.67, 6.67, 7.95, 0.34, "GPT-5.6 Terra · OpenAI Responses API · evidence-grounded RAG · SQLite ticket database · Streamlit · human-controlled release", 6.2, TEXT)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.84), Inches(6.48), Inches(0.01), Inches(0.53))
    divider.fill.solid(); divider.fill.fore_color.rgb = LINE; divider.line.fill.background()
    textbox(slide, 9.0, 6.48, 3.9, 0.18, "STUDENT INFORMATION", 6.2, NAVY, True)
    textbox(slide, 9.0, 6.67, 3.9, 0.33, "Ikram Jemlaoui · AI Language Models and Business Applications · 16 September 2026", 6.1, TEXT)
    textbox(slide, 0.24, 7.23, 12.85, 0.14, "Sources: project implementation; OpenAI API documentation; NIST AI RMF; OWASP GenAI Top 10. Impact figures are pilot targets, not measured production results.", 4.8, MUTED)

    prs.core_properties.title = "AegisDesk AI — Six-Section Agentic Application One-Pager"
    prs.core_properties.subject = "LLM & Business Applications assignment"
    prs.core_properties.author = "Ikram Jemlaoui"
    prs.core_properties.comments = "Six-section design with agentic approach and expected business impact."
    prs.save(OUTPUT)
    print(f"Generated: {OUTPUT}")
    print(f"Slides: {len(prs.slides)}; shapes: {len(slide.shapes)}")


if __name__ == "__main__":
    build()
