"""Correct only page 1 of the v6 architecture report.

The source document is preserved. The corrected copy keeps the existing layout,
styles, tables, and pages 2-5 unchanged while replacing the four repeated
overview sections with the required Business Layer content.
"""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "AegisDesk_AI_Architecture_5_Pages_v6.docx"
OUTPUT = ROOT / "deliverables" / "AegisDesk_AI_Architecture_5_Pages_v7.docx"


REPLACEMENTS = {
    1: "1. Business Objectives",
    2: (
        "Create one intelligent request channel for the whole company; improve request completeness "
        "before submission; route each request to the correct accountable department; prioritise urgent "
        "and high-impact work; reuse approved knowledge and verified resolutions; and keep people "
        "responsible for final decisions."
    ),
    4: "2. Key Performance Indicators (KPIs)",
    5: "Request completeness at first submission and the number of clarification exchanges required.",
    6: "First-time routing accuracy and the rate of transfers or reassignments between departments.",
    7: "Time to first useful response, total resolution time, SLA compliance and urgent-request response time.",
    8: "Self-service resolution, suggested-answer acceptance, employee satisfaction and verified knowledge reuse.",
    9: "3. Stakeholders",
    10: "Requesters: employees and teams from every company department.",
    11: "Resolver teams: IT, Data, Operations, Finance, HR, Security, Ecommerce and other specialist departments.",
    12: "Operational owners: department managers, service and process owners, and company knowledge owners.",
    13: "Control and platform stakeholders: security, privacy/legal, IT and AI platform operations, and the executive sponsor.",
    14: "4. Expected Value",
    15: (
        "Employees receive faster guidance with fewer handoffs. Resolver teams receive complete, prioritised "
        "and evidence-backed requests. Managers gain clear ownership and service visibility. The company "
        "reduces repeated work, accelerates resolution and turns successful outcomes into reusable operational knowledge."
    ),
}


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    document = Document(SOURCE)
    if len(document.paragraphs) != 87 or len(document.tables) != 20:
        raise RuntimeError("Unexpected source structure; refusing to modify the report.")

    original_after_page_one = [paragraph.text for paragraph in document.paragraphs[18:]]
    original_tables = [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]

    for paragraph_index, new_text in REPLACEMENTS.items():
        paragraph = document.paragraphs[paragraph_index]
        if not paragraph.runs:
            paragraph.add_run(new_text)
            continue
        # Reuse the existing first run so its font, size, weight and colour remain unchanged.
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""

    document.core_properties.title = "AegisDesk AI Architecture — Five-Page Report"
    document.core_properties.subject = "Corrected Business Layer: objectives, KPIs, stakeholders and expected value"
    document.save(OUTPUT)

    verification = Document(OUTPUT)
    if [paragraph.text for paragraph in verification.paragraphs[18:]] != original_after_page_one:
        raise RuntimeError("Pages 2-5 changed unexpectedly.")
    verified_tables = [[cell.text for row in table.rows for cell in row.cells] for table in verification.tables]
    if verified_tables != original_tables:
        raise RuntimeError("A table changed unexpectedly.")

    print(f"Generated: {OUTPUT}")
    print(f"Paragraphs: {len(verification.paragraphs)}; tables: {len(verification.tables)}")
    print("Verified: all tables and pages 2-5 content remain unchanged")


if __name__ == "__main__":
    main()
