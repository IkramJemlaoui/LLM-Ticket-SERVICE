"""Generate the assignment's required one-slide PPTX and five-page DOCX.

The files are intentionally generated from code so content and visual labels stay
version-controlled and reproducible.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor as DocxRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape as rl_landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"
AUTHOR = "Ikram Jemlaoui"
PROGRAM = "AI Language Models and Business Applications"
DATE_LABEL = "15 September 2026"


NAVY = RGBColor(20, 50, 79)
PANEL = RGBColor(255, 255, 255)
PANEL_2 = RGBColor(235, 247, 255)
WHITE = RGBColor(244, 248, 255)
MUTED = RGBColor(72, 97, 121)
MINT = RGBColor(11, 143, 145)
VIOLET = RGBColor(139, 92, 246)
BLUE = RGBColor(96, 165, 250)
AMBER = RGBColor(245, 158, 11)
RED = RGBColor(251, 113, 133)


def ppt_textbox(slide, x, y, w, h, text, size=12, color=WHITE, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = PptInches(0.04)
    frame.margin_top = frame.margin_bottom = PptInches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = PptPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def ppt_card(slide, x, y, w, h, number, title, body, accent):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PptInches(x), PptInches(y), PptInches(w), PptInches(h),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = RGBColor(171, 204, 226)
    card.line.width = PptPt(1.0)

    marker = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PptInches(x + 0.18), PptInches(y + 0.16), PptInches(0.38), PptInches(0.34),
    )
    marker.fill.solid()
    marker.fill.fore_color.rgb = accent
    marker.line.fill.background()
    marker.text_frame.clear()
    p = marker.text_frame.paragraphs[0]
    p.text = str(number)
    p.font.name = "Aptos Display"
    p.font.size = PptPt(11)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    ppt_textbox(slide, x + 0.66, y + 0.17, w - 0.83, 0.3, title.upper(), 11, accent, True, "Aptos Display")
    ppt_textbox(slide, x + 0.2, y + 0.6, w - 0.4, h - 0.72, body, 9.2, MUTED, False)


def build_one_pager(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(238, 248, 255)

    # Restrained visual accents keep the one-pager light and readable.
    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, PptInches(10.9), PptInches(-1.25), PptInches(3.4), PptInches(3.4))
    glow.fill.solid(); glow.fill.fore_color.rgb = RGBColor(144, 221, 232); glow.fill.transparency = 45
    glow.line.fill.background()
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(0.1), prs.slide_height)
    band.fill.solid(); band.fill.fore_color.rgb = MINT; band.line.fill.background()

    ppt_textbox(slide, 0.35, 0.2, 8.8, 0.28, "COMPANY-WIDE REQUEST INTELLIGENCE · ONE-PAGER", 9, MINT, True)
    ppt_textbox(slide, 0.35, 0.51, 8.6, 0.52, "AegisDesk AI", 26, NAVY, True, "Aptos Display")
    ppt_textbox(
        slide, 0.35, 1.02, 10.2, 0.42,
        "One front door for every internal request. The right context, the right answer and the right team—before delays begin.",
        11.2, MUTED,
    )
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(10.95), PptInches(0.44), PptInches(1.92), PptInches(0.54))
    pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor(220, 247, 245); pill.line.color.rgb = MINT
    pill.text_frame.clear(); p = pill.text_frame.paragraphs[0]; p.text = "REAL LLM · HUMAN CONTROL"; p.font.size = PptPt(8.5); p.font.bold = True; p.font.color.rgb = MINT; p.alignment = PP_ALIGN.CENTER
    pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    gap = 0.16
    left = 0.35
    width = (12.63 - gap) / 2
    y1, y2 = 1.58, 4.08
    h = 2.25
    ppt_card(
        slide, left, y1, width, h, 1, "Application overview",
        "AegisDesk is a cross-department request platform for the whole company. Any team can ask another team for support, clarification, correction or new work. Coordinated AI agents turn plain-language requests into complete, prioritised and correctly routed cases, then bring forward approved guidance and previously verified solutions.",
        BLUE,
    )
    ppt_card(
        slide, left + width + gap, y1, width, h, 2, "Key Challenges Addressed",
        "01  Incomplete requests create clarification loops.\n02  Wrong routing delays the responsible department.\n03  Useful answers are spread across documents, cases and people.\n04  Urgent, high-impact work can disappear in a shared queue.\n05  Unsupported AI answers create trust, privacy and safety risks.",
        VIOLET,
    )
    ppt_card(
        slide, left, y2, width, h, 3, "Core capabilities (services delivered)",
        "EMPLOYEE PORTAL  Clarifies requests before submission and offers safe self-service.\n\nSMART RESOLUTION  Classifies, prioritises and recommends the accountable team.\n\nGROUNDED HELP  Searches approved knowledge and verified cases, drafts a cited answer and asks for missing facts.\n\nSUPPORT WORKSPACE  Lets specialists edit, approve, reject, reassign, escalate or resolve.",
        MINT,
    )
    ppt_card(
        slide, left + width + gap, y2, width, h, 4, "Differentiation",
        "AGENTIC DESIGN  Five role-separated agents collaborate through a bounded orchestrator.\n\nVALUE FOCUS  Fewer clarification loops, faster response, better routing and reusable knowledge.\n\nINTEGRATION  Designed for company knowledge, case-management and collaboration systems.\n\nSCALABILITY  Permission-aware connectors and observable operations support growth. Humans approve consequential actions.",
        AMBER,
    )
    ppt_textbox(slide, 0.35, 6.55, 9.9, 0.28, "OpenAI Responses API · evidence-grounded RAG · five bounded agents · permission-aware integration design · human-controlled release", 7.7, MUTED)
    ppt_textbox(slide, 10.25, 6.48, 2.6, 0.55, f"{AUTHOR}\n{PROGRAM} · {DATE_LABEL}", 7.5, MUTED, False, align=PP_ALIGN.RIGHT)
    ppt_textbox(slide, 0.35, 7.11, 12.5, 0.18, "Sources: project implementation; official OpenAI API documentation; NIST AI RMF; OWASP GenAI Top 10. Pilot impact is measured after deployment.", 5.8, RGBColor(83, 106, 130))

    prs.core_properties.title = "AegisDesk AI — Agentic Application One-Pager"
    prs.core_properties.subject = "LLM & Business Applications assignment"
    prs.core_properties.author = AUTHOR
    prs.core_properties.comments = "Generated from version-controlled project content; all pilot figures are targets."
    prs.save(path)


DOC_NAVY = "14324F"
DOC_BLUE = "12345A"
DOC_MINT = "16A394"
DOC_PALE = "EAF7F5"
DOC_LIGHT = "EEF8FF"
DOC_WHITE = "FFFFFF"
DOC_MUTED = DocxRGB(73, 91, 115)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold=False, color="07111F", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = DocxRGB.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AEGISDESK AI  ·  ")
    run.font.name = "Aptos"; run.font.size = Pt(8); run.font.color.rgb = DOC_MUTED
    fld_char_1 = OxmlElement("w:fldChar"); fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText"); instr_text.set(qn("xml:space"), "preserve"); instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar"); fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def add_title_band(doc: Document, eyebrow: str, title: str, subtitle: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    table.style = "Table Grid"
    shade_cell(cell, DOC_LIGHT)
    cell.margin_top = Cm(0.35); cell.margin_bottom = Cm(0.35)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(eyebrow.upper()); r.font.name = "Aptos"; r.font.size = Pt(8); r.bold = True; r.font.color.rgb = DocxRGB.from_string(DOC_MINT)
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(3)
    r = p2.add_run(title); r.font.name = "Aptos Display"; r.font.size = Pt(20); r.bold = True; r.font.color.rgb = DocxRGB.from_string(DOC_NAVY)
    p3 = cell.add_paragraph(); p3.paragraph_format.space_after = Pt(0)
    r = p3.add_run(subtitle); r.font.name = "Aptos"; r.font.size = Pt(9); r.font.color.rgb = DOC_MUTED
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level=1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.02
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(1.5)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], DOC_BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=DOC_WHITE, size=8)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            shade_cell(cells[index], DOC_WHITE if row_index % 2 == 0 else DOC_LIGHT)
            set_cell_text(cells[index], value, size=7.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, label: str, text: str, color=DOC_PALE) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label.upper()}  "); r.bold = True; r.font.color.rgb = DocxRGB.from_string(DOC_MINT)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = DocxRGB.from_string(DOC_NAVY)
    for name, size, color in [("Title", 25, DOC_NAVY), ("Heading 1", 14, DOC_BLUE), ("Heading 2", 10.5, DOC_MINT)]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = DocxRGB.from_string(color)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    properties = doc.core_properties
    properties.title = "AegisDesk AI — Architecture Document"
    properties.subject = "LLM & Business Applications — six-layer agentic architecture"
    properties.author = AUTHOR
    properties.keywords = "LLM, multi-agent, service desk, governance, RAG"
    properties.comments = "Five-page assignment document generated from version-controlled project content."


def _build_architecture_document_legacy(path: Path) -> None:
    doc = Document()
    configure_doc(doc)

    # Page 1 — Business layer.
    add_title_band(doc, "Page 1 · Executive overview & business layer", "AegisDesk AI", "An evidence-first, human-accountable multi-agent copilot for IT service operations")
    add_heading(doc, "Application overview")
    add_body(doc, "AegisDesk serves an ecommerce company's IT, Ecommerce Platform, Data Engineering, and BI & Analytics teams. It converts workplace incidents, checkout/order problems, pipeline failures, and Power BI/KPI questions into routed recommendations with a visible similar-case solution and approved diagnostic guidance. Five specialist agents operate under deterministic orchestration and human approval.")
    add_heading(doc, "Business problem and value", 2)
    add_bullets(doc, [
        "Incomplete tickets create repeated clarification and slow first response.",
        "Scattered knowledge produces inconsistent troubleshooting and avoidable rework.",
        "Opaque AI drafts introduce hallucination, privacy, and automation-bias risk.",
        "AegisDesk makes evidence and control visible while reducing repetitive reading, searching, and drafting.",
    ])
    add_heading(doc, "Users, stakeholders and impacted process", 2)
    add_body(doc, "Primary users: IT support, Ecommerce Platform, Data Engineering, BI & Analytics, and service-desk managers. Beneficiaries are employees and ecommerce operations. Proposed sponsors: Head of IT Support and Ecommerce Technology lead. Control stakeholders include security, privacy/legal, business KPI owners, and knowledge owners.")
    add_callout(doc, "Target process", "Ticket intake → safety screen → triage → evidence search → draft → quality validation → human approval or escalation.")
    add_heading(doc, "Pilot value scorecard", 2)
    add_table(doc, ["Metric", "4-week target", "Evidence"], [
        ["Time to first qualified draft", "−40% vs baseline", "Open-to-approval timestamps"],
        ["Triage rework", "−25%", "Supervisor correction audit"],
        ["Evidence coverage", "≥85% covered intents", "Workflow telemetry"],
        ["Acceptance without material rewrite", "≥60%", "Approval/edit events"],
        ["Unsafe autonomous sends", "0", "Release-control audit"],
    ])
    add_body(doc, "ROI logic: minutes saved × eligible ticket volume × loaded labor cost, minus model, platform, integration, training, and governance costs. All figures above are pilot targets—not measured production results.")
    doc.add_page_break()

    # Page 2 — Agentic/application layer.
    add_title_band(doc, "Page 2 · Agentic / application layer", "Five roles. One bounded workflow.", "Role separation makes each decision testable, observable and governable")
    add_heading(doc, "Agent ecosystem")
    add_table(doc, ["Agent", "Input → output", "Authority / boundary"], [
        ["Safety & Privacy", "Raw ticket → sanitized ticket or stop", "Redact, truncate, stop; no generation"],
        ["Triage", "Sanitized ticket → summary, gaps, category, priority, confidence", "Role-specific LLM/mock; fixed taxonomy"],
        ["Knowledge", "Ticket + triage → approved passages and scores", "Read-only KB; threshold; unique top-3"],
        ["Resolution", "Triage + evidence → draft and citations", "Evidence-only LLM/mock; no send"],
        ["Risk & Quality", "Decisions + draft → pass, block or escalate", "Deterministic validation; no self-approval"],
    ])
    add_heading(doc, "Orchestration and decision logic", 2)
    flow = doc.add_table(rows=1, cols=6)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"]):
        shade_cell(flow.cell(0, index), DOC_MINT if index in {0, 5} else DOC_BLUE)
        set_cell_text(flow.cell(0, index), value, bold=True, color=DOC_WHITE, size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_bullets(doc, [
        "STOP: prompt-injection indicators end downstream processing and require security review.",
        "ESCALATE: no evidence above the relevance threshold routes to a specialist.",
        "AWAIT INFORMATION: three or more material gaps produce clarification questions.",
        "HUMAN REVIEW: every valid draft waits for approval; high/critical impact has an explicit reason.",
        "BOUND: maximum six trace steps, no arbitrary tools/web, no autonomous ticket change/send; local resolution publishing requires a human.",
    ])
    add_heading(doc, "Human-in-the-loop", 2)
    add_body(doc, "Page 1, the Employee Help Portal, analyses before submission and shows safe self-service, a similar solved case, documentation, missing facts, classification, and recommended team. Page 2, the Support Workspace, shows queue/SLA, evidence, editable final response, quality, trace, and approve/reject/reassign/escalate controls. Resolution publishing remains a separate human action. Telemetry excludes draft text.")
    add_callout(doc, "Differentiator", "AegisDesk does not hide agent reasoning behind one chat response: it exposes bounded handoffs, accepted evidence, stop/escalate logic, and accountable release.")
    doc.add_page_break()

    # Page 3 — LLM and data.
    add_title_band(doc, "Page 3 · LLM & data layers", "Language intelligence grounded in approved evidence", "Use an LLM where language adds value; keep critical controls deterministic")
    add_heading(doc, "LLM layer")
    add_body(doc, "Two narrow model roles are supported. The Triage Agent returns strict structured fields using an allowed taxonomy and a no-invention instruction. The Resolution Agent receives the triage decision and only accepted article/case evidence, then returns a concise draft and confidence note. Safety, retrieval, validation, orchestration, and approval are not delegated to the model.")
    add_table(doc, ["Decision", "Current PoC", "Selection/control rationale"], [
        ["Provider", "Deterministic mock; optional Azure/OpenAI-compatible", "Reproducible by default; managed path is swappable"],
        ["Prompting", "Two role-specific strict-JSON prompts", "Separation of triage and drafting responsibilities"],
        ["Parameters", "Temperature ≤0.4; output ≤800 tokens", "Stable, concise, cost-bounded output"],
        ["Failure", "25-second timeout + mock fallback", "Continue safely and expose fallback"],
        ["Evaluation", "Regression/safety suite; pilot benchmark planned", "Model chosen on accuracy, latency, privacy and cost"],
    ])
    add_heading(doc, "Data layer and retrieval", 2)
    add_body(doc, "Runtime tickets, activity, article metadata, and verified resolutions use local SQLite. Reproducible sources are 27 ticket fixtures (including 16 verified fixes), 20 approved Markdown articles, and 15 governed routing evaluations spanning workplace IT, ecommerce, data, Finance, HR, and cross-department support. Tickets are restricted operational data; pilot content requires named ownership, access, retention, and freshness controls.")
    add_callout(doc, "RAG flow", "Approved article paragraphs + human-verified cases → TF-IDF matrix → triage-enriched query → cosine similarity → score ≥0.08 → unique top-3 → grounded draft + typed citations. Retrieval supplies context; it does not train the LLM.")
    add_bullets(doc, [
        "No open-web retrieval and no vector database are used in the current PoC.",
        "Subject, description and requester are sanitized before optional cloud inference.",
        "Low-scoring evidence is rejected rather than presented as reliable grounding.",
        "Production additions: RBAC, encryption, owner/version metadata, freshness SLA, retention/deletion, residency and ingestion validation.",
        "Embeddings/hybrid retrieval are adopted only if a labelled benchmark demonstrates material improvement.",
    ])
    doc.add_page_break()

    # Page 4 — GRC.
    add_title_band(doc, "Page 4 · Governance, risk & compliance", "Trust is a workflow, not a slogan", "Technical controls, ownership, evidence and residual-risk decisions")
    add_heading(doc, "Risk and control register")
    add_table(doc, ["Risk", "Implemented control", "Residual / pilot action"], [
        ["Prompt injection", "Pre-agent detection and immediate stop", "Adversarial evaluation; indirect-injection controls"],
        ["Sensitive disclosure", "Field-level PII redaction; no raw prompt/draft logs", "DPIA, DLP, retention/residency and incident process"],
        ["Hallucination", "Approved KB, threshold, citations, independent quality check", "Groundedness benchmark and sampled QA"],
        ["Wrong triage", "Fixed taxonomy, confidence, editable recommendation", "Labelled accuracy and drift monitoring"],
        ["Unsafe output", "Prohibited-language and citation-provenance validation", "Expanded classifier and red-team suite"],
        ["Automation bias", "Visible trace/evidence; mandatory human approval", "Training and override/material-edit audit"],
        ["Provider failure", "Timeout, fallback and visible status", "Circuit breaker, SLO and bounded retry policy"],
    ])
    add_heading(doc, "Governance operating model", 2)
    add_bullets(doc, [
        "Accountable owner: service-desk product owner; co-owners: security, privacy/legal, knowledge and platform operations.",
        "Weekly pilot review of safety, quality, adoption, latency and cost; monthly release gate.",
        "Prompt/model/config changes require version, benchmark, sign-off, monitoring threshold and rollback criterion.",
        "Incidents and exceptions record impact, owner, containment, correction and prevention action.",
        "Priority uses operational impact—not seniority or sensitive characteristics; evaluation checks channel/language/accessibility disparities.",
    ])
    add_heading(doc, "Compliance posture", 2)
    add_body(doc, "The design applies privacy-by-design and data-minimization principles, but this classroom PoC does not claim GDPR, AI Act, ISO, SOC 2, or other certification. A real pilot requires organizational legal assessment, vendor terms, processing records, access/retention decisions, and security testing.")
    add_callout(doc, "Release rule", "The Risk & Quality Agent may release a draft only to human consideration. It can never authorize customer delivery on its own.")
    doc.add_page_break()

    # Page 5 — Operations, roadmap and sources.
    add_title_band(doc, "Page 5 · Operations, monitoring & references", "Measure value, quality, safety and cost together", "Observe every handoff without logging the customer's raw content")
    add_heading(doc, "Operational observability")
    add_table(doc, ["Dimension", "Signals", "Pilot response"], [
        ["Reliability", "Success, fallback, p95 latency, availability", "Alert if fallback >5% or p95 >10 s"],
        ["Quality", "Triage accuracy, evidence precision, groundedness, edits", "Review failed/edited samples weekly"],
        ["Safety", "Refusal, safety review, PII events, unauthorized release", "Immediate incident for PII or release breach"],
        ["Cost", "Prompt/completion tokens and cost/workflow", "Model/prompt optimization at release gate"],
        ["Business", "Draft time, rework, SLA, acceptance, satisfaction", "Continue/stop/scale at four-week gate"],
    ])
    add_heading(doc, "Continuous-improvement loop", 2)
    add_body(doc, "Collect privacy-minimized metrics → review escalations and material edits → label root cause → change one versioned component → run regression and adversarial evaluation → approve/reject release → monitor pilot → roll back if thresholds fail.")
    add_heading(doc, "Feasible rollout", 2)
    add_bullets(doc, [
        "Week 0: governance sign-off, baseline, labelled evaluation set and agent training.",
        "Weeks 1–2: shadow mode on one queue; compare suggestions without operational release.",
        "Weeks 3–4: assisted mode with human approval; measure target KPIs and incidents.",
        "Gate: stop, refine, or scale based on evidence—not enthusiasm. Roadmap: SSO/RBAC, ITSM API, managed monitoring and evaluated hybrid retrieval.",
    ])
    add_heading(doc, "References and disclosure", 2)
    sources = [
        "Assignment — LLM & Business Applications: Designing an Agentic LLM-Based Application (course-provided DOCX and templates).",
        "NIST AI Risk Management Framework and Generative AI Profile — https://www.nist.gov/itl/ai-risk-management-framework",
        "OWASP GenAI Security Project, Top 10 — https://genai.owasp.org/llm-top-10/",
        "EU General Data Protection Regulation 2016/679 — https://eur-lex.europa.eu/eli/reg/2016/679/",
        "scikit-learn TfidfVectorizer — https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html",
        "Streamlit documentation — https://docs.streamlit.io/",
        "Microsoft Azure OpenAI/Foundry REST reference — https://learn.microsoft.com/en-us/azure/foundry/openai/latest",
        "Microsoft Learn, Power BI KPI, refresh, and star-schema guidance — https://learn.microsoft.com/en-us/power-bi/visuals/power-bi-visualization-kpi",
        "Stripe Documentation, Checkout fulfillment and Payment Intents — https://docs.stripe.com/checkout/fulfillment",
    ]
    add_bullets(doc, sources)
    add_callout(doc, "AI/tool acknowledgement", "Generative AI assisted with code, document structure and wording. The student reviewed, tested and remains responsible for all claims. Python, Streamlit, scikit-learn, python-docx and python-pptx were used. Tickets and KB content are synthetic.")
    add_body(doc, f"Prepared by {AUTHOR} · {PROGRAM} · {DATE_LABEL}")

    doc.save(path)


def build_architecture_document(path: Path) -> None:
    """Build the five-page architecture report using the agreed light visual system."""
    doc = Document()
    configure_doc(doc)

    # Page 1 — required one-pager sections, in the assignment's exact order.
    add_title_band(
        doc,
        "Page 1 · Business layer",
        "AegisDesk AI",
        "One intelligent front door for every internal company request",
    )
    add_heading(doc, "1. Application Overview")
    add_body(
        doc,
        "AegisDesk is a company-wide request intelligence platform. Any department can ask another department for support, clarification, correction, or new work. Employees explain the need in plain language; coordinated AI agents clarify it, assess impact, recommend the accountable team, retrieve approved company knowledge and verified resolutions, and prepare the next best action. Requesters receive faster guidance, resolver teams receive complete cases, and the company turns successful resolutions into reusable knowledge.",
    )
    add_callout(
        doc,
        "Product promise",
        "The right context, the right answer and the right team—before delays begin.",
    )
    add_heading(doc, "2. Key Challenges Addressed")
    add_bullets(doc, [
        "Incomplete requests force teams into repeated clarification loops.",
        "Wrong routing delays ownership and can hide urgent work in the wrong queue.",
        "Answers are spread across documents, past cases and individual employees, so work is repeatedly rediscovered.",
        "AI answers without visible sources or controls may be incorrect, unsafe, or trusted too quickly.",
    ])
    add_heading(doc, "3. Core Capabilities (Services Delivered)")
    add_bullets(doc, [
        "Employee Request Portal: clarifies the request before submission, asks for missing facts and offers safe self-service when reliable help exists.",
        "Smart classification and routing: identifies request type, urgency, business impact and the accountable destination team across departments.",
        "Evidence-grounded assistance: searches approved knowledge and verified resolutions, shows sources and prepares a response for review.",
        "Support Workspace: prioritised queue, SLA, editable answer, evidence, confidence, agent trace, reassignment, escalation and approval controls.",
    ])
    add_heading(doc, "4. Differentiation")
    add_body(
        doc,
        "AegisDesk combines five role-separated agents, deterministic orchestration, real LLM language intelligence and governed RAG. It is designed to connect with enterprise knowledge and case systems, scale across departments, and measure routing quality, speed, reuse and safety. Human approval remains mandatory for consequential actions.",
    )
    doc.add_page_break()

    # Page 2 — Agentic/Application layer.
    add_title_band(
        doc,
        "Page 2 · Agentic / Application layer",
        "Coordinated intelligence with clear responsibility",
        "Five specialist agents collaborate; the orchestrator controls every handoff",
    )
    add_heading(doc, "Controlled workflow")
    flow = doc.add_table(rows=1, cols=6)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow.style = "Table Grid"
    for index, value in enumerate(["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"]):
        shade_cell(flow.cell(0, index), DOC_PALE if index in {0, 5} else DOC_LIGHT)
        set_cell_text(flow.cell(0, index), value, bold=True, color=DOC_BLUE, size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_heading(doc, "Five separated responsibilities", 2)
    add_callout(doc, "Safety & Privacy Agent", "Screens input, minimises personal data and stops prompt attacks before cloud inference.")
    add_callout(doc, "Triage Agent", "Uses the live LLM to summarise, identify missing facts and classify the request; approved policy maps the category to a team.", DOC_LIGHT)
    add_callout(doc, "Knowledge Agent", "Searches approved company documentation and human-verified resolutions; weak matches are rejected.")
    add_callout(doc, "Resolution Agent", "Uses the live LLM to draft a response from retrieved evidence only; it cannot send or execute actions.", DOC_LIGHT)
    add_callout(doc, "Risk & Quality Agent", "Checks taxonomy, evidence provenance, unsafe language and quality before human review.")
    add_heading(doc, "Orchestration and controlled autonomy", 2)
    add_bullets(doc, [
        "STOP when unsafe or instruction-injection language is detected.",
        "ASK when material information is missing; ESCALATE when no approved evidence is strong enough.",
        "ROUTE through an approved cross-department taxonomy; prioritise by impact and urgency.",
        "RELEASE only to a human decision point: approve, edit, reject, reassign, escalate, or resolve.",
        "BOUND autonomy to six visible steps, no open-web search, no arbitrary tools and no autonomous external message.",
    ])
    add_heading(doc, "Two role-based experiences", 2)
    add_body(doc, "The Employee Request Portal serves every department as a requester. The Support Workspace serves whichever department owns the request. A team can therefore be a requester in one workflow and a resolver in another.")
    doc.add_page_break()

    # Page 3 — LLM and Data layers.
    add_title_band(
        doc,
        "Page 3 · LLM and Data layers",
        "Real language intelligence, grounded in company evidence",
        "The LLM interprets and drafts; policy, retrieval and release remain controlled",
    )
    add_heading(doc, "LLM layer")
    add_body(
        doc,
        "The application now defaults to a real OpenAI-compatible deployment. The OpenAI path calls the Responses API with the configured model, strict JSON Schema outputs, stateless storage disabled, a bounded output budget and a request timeout. The Triage and Resolution agents are the only model-backed roles. Live mode fails visibly when credentials or the provider are unavailable; it does not silently present simulated results.",
    )
    add_bullets(doc, [
        "Triage prompt: understand technical, operational, data, finance, people and business requests; return summary, missing facts, category, priority and confidence.",
        "Resolution prompt: use only accepted internal evidence; distinguish an analogous past case from a guaranteed solution; ask or escalate instead of inventing.",
        "Structured outputs: allowed categories and priorities are schema-constrained; routing remains application policy.",
        "Privacy: identifiers are redacted before provider calls and raw request/draft text is excluded from telemetry.",
    ])
    add_callout(doc, "Configuration", "LLM_PROVIDER=openai · OPENAI_BASE_URL=https://api.openai.com/v1 · OPENAI_MODEL=gpt-5.6-terra · secret supplied only through OPENAI_API_KEY")
    add_heading(doc, "Data and RAG layer")
    add_body(
        doc,
        "In operation, authorised sources can include company knowledge pages such as Confluence, existing service records and human-verified resolutions. Retrieval filters by permissions, approval status, ownership and freshness; the Resolution Agent receives only the accepted excerpts and their source metadata.",
    )
    add_callout(doc, "RAG sequence", "Request + triage context  →  permission-aware search  →  relevance threshold  →  top evidence  →  cited draft  →  quality validation. Retrieval supplies current context; it does not retrain the LLM.", DOC_LIGHT)
    add_bullets(doc, [
        "Prototype storage: local SQLite for requests, activities, article metadata, retrieval feedback and verified resolutions.",
        "Prototype retrieval: transparent TF-IDF/cosine scoring over the local demonstration corpus.",
        "Production evolution: enterprise database, SSO/RBAC, encrypted storage and evaluated hybrid/vector retrieval when benchmarks justify it.",
    ])
    doc.add_page_break()

    # Page 4 — Governance, Risk and Compliance layer.
    add_title_band(
        doc,
        "Page 4 · Governance, Risk and Compliance layer",
        "Trust is built into the workflow",
        "Visible evidence, bounded authority, accountable owners and measurable controls",
    )
    add_heading(doc, "Risk-to-control design")
    add_callout(doc, "Incorrect or invented answer", "Approved evidence, source display, retrieval threshold, structured output, independent validation and human approval.")
    add_callout(doc, "Sensitive information", "Input minimisation, PII redaction before cloud inference, secret isolation, permission-aware retrieval and no raw-content telemetry.", DOC_LIGHT)
    add_callout(doc, "Prompt or document injection", "Pre-agent screening, no open web or arbitrary tools, untrusted attachment handling and adversarial evaluation.")
    add_callout(doc, "Wrong team or priority", "Fixed taxonomy, deterministic routing map, impact-based priority, visible confidence and easy human reassignment.", DOC_LIGHT)
    add_callout(doc, "Automation bias", "The UI shows evidence and uncertainty; specialists can edit, reject or escalate and AI cannot approve itself.")
    add_heading(doc, "Governance operating model", 2)
    add_bullets(doc, [
        "Product owner accountable for outcomes; knowledge owners accountable for source accuracy and review dates.",
        "Security, privacy/legal and platform owners approve provider, access, retention, residency and incident controls.",
        "Every prompt, model, taxonomy and threshold change is versioned, evaluated, approved and reversible.",
        "Priority uses operational impact rather than seniority or sensitive personal characteristics.",
        "The design supports privacy and responsible-AI principles but does not claim legal or security certification without organisational assessment.",
    ])
    add_heading(doc, "Human accountability", 2)
    add_body(doc, "The system may recommend and prepare. Only an authorised person may release a consequential response, reassign ownership, close a case or publish a resolution into reusable knowledge.")
    doc.add_page_break()

    # Page 5 — Operations/Monitoring plus all supplementary material.
    add_title_band(
        doc,
        "Page 5 · Operations, Monitoring and References",
        "Operate for measurable value",
        "Quality, reliability, cost and safety are monitored together",
    )
    add_heading(doc, "Business and service measures")
    add_bullets(doc, [
        "Request completeness and first-time routing accuracy.",
        "Time to first useful response, resolution time, SLA compliance and urgent-case response.",
        "Self-service resolution, suggested-answer acceptance and reduction in repeated investigations.",
        "Employee satisfaction, specialist edits/overrides and verified knowledge reuse.",
    ])
    add_heading(doc, "LLM and operational monitoring", 2)
    add_bullets(doc, [
        "Provider availability, end-to-end and per-agent latency, timeout/failure rate and token usage.",
        "Retrieval relevance, evidence coverage, grounding quality, safety stops, PII events and unauthorised-release incidents.",
        "Continuous improvement: review outcomes → label root cause → change one versioned component → regression and adversarial evaluation → approval → monitored release or rollback.",
    ])
    add_heading(doc, "Integration and scale", 2)
    add_body(doc, "A production deployment can connect to Confluence or another approved knowledge platform, an existing case-management system, identity/SSO and collaboration channels. Permission checks, source ownership and freshness travel with indexed content. Stateless application instances, background ingestion, managed storage and queue-based processing support growth across departments.")
    add_heading(doc, "References and implementation disclosure", 2)
    add_bullets(doc, [
        "Course assignment: LLM & Business Applications — Designing an Agentic LLM-Based Application.",
        "OpenAI Responses API — https://developers.openai.com/api/reference/python/resources/responses/methods/create",
        "Atlassian Confluence search and permissions APIs — https://developer.atlassian.com/cloud/confluence/rest/v1/",
        "NIST AI Risk Management Framework — https://www.nist.gov/itl/ai-risk-management-framework",
        "OWASP GenAI Security Project — https://genai.owasp.org/llm-top-10/",
        "Streamlit — https://docs.streamlit.io/ · scikit-learn TF-IDF — https://scikit-learn.org/stable/modules/feature_extraction.html",
    ])
    add_callout(doc, "Disclosure", "Generative AI assisted implementation and document drafting; the student reviewed the content and remains responsible for all claims. The repository includes local demonstration data. Production operation requires authorised company sources, credentials, security review and measured results.")
    add_body(doc, f"Prepared by {AUTHOR} · {PROGRAM} · {DATE_LABEL}")
    doc.save(path)


def _build_architecture_pdf_legacy(path: Path) -> None:
    """Build a fixed five-page PDF directly, without Microsoft Office automation."""
    base = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle("AegisBody", parent=base["BodyText"], fontName="Helvetica", fontSize=8.6, leading=11.2, textColor=colors.HexColor("#172B45"), spaceAfter=5),
        "h1": ParagraphStyle("AegisH1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#12345A"), spaceBefore=5, spaceAfter=5),
        "h2": ParagraphStyle("AegisH2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#168E83"), spaceBefore=5, spaceAfter=3),
        "small": ParagraphStyle("AegisSmall", parent=base["BodyText"], fontName="Helvetica", fontSize=7.1, leading=9, textColor=colors.HexColor("#536A85"), spaceAfter=2),
        "bullet": ParagraphStyle("AegisBullet", parent=base["BodyText"], fontName="Helvetica", fontSize=8.2, leading=10.5, leftIndent=11, firstLineIndent=-7, textColor=colors.HexColor("#172B45"), spaceAfter=2),
    }

    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=1.35 * cm, leftMargin=1.35 * cm,
        topMargin=1.25 * cm, bottomMargin=1.2 * cm, title="AegisDesk AI - Architecture Document",
        author=AUTHOR, subject="Five-page six-layer agentic architecture", pageCompression=0,
    )

    def header_footer(canvas, doc):
        canvas.saveState()
        width, height = A4
        canvas.setFillColor(colors.HexColor("#07111F"))
        canvas.rect(0, height - 0.46 * cm, width, 0.46 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#16A394"))
        canvas.rect(0, height - 0.46 * cm, 1.15 * cm, 0.46 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#62758D"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(1.35 * cm, 0.58 * cm, "AEGISDESK AI  |  IMPLEMENTED, TESTED AND HUMAN-CONTROLLED")
        canvas.drawRightString(width - 1.35 * cm, 0.58 * cm, f"PAGE {doc.page}")
        canvas.restoreState()

    def title_band(kicker, title, subtitle):
        content = Paragraph(
            f'<font color="#39D8C3" size="8"><b>{kicker.upper()}</b></font><br/>'
            f'<font color="#FFFFFF" size="20"><b>{title}</b></font><br/>'
            f'<font color="#B7C6D9" size="8.5">{subtitle}</font>',
            ParagraphStyle("band", parent=styles["body"], leading=16, spaceAfter=0),
        )
        table = Table([[content]], colWidths=[17.9 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#07111F")),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#12345A")),
            ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]))
        return [table, Spacer(1, 6)]

    def p(text, style="body"):
        return Paragraph(text, styles[style])

    def bullets(items):
        return [p(f"&#8226;&nbsp; {item}", "bullet") for item in items]

    def table(headers, rows, widths=None):
        data = [[p(f"<b>{item}</b>", "small") for item in headers]]
        data.extend([[p(item, "small") for item in row] for row in rows])
        element = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12345A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("BACKGROUND", (0, 1), (-1, -1), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D5E3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return element

    def callout(label, text):
        element = Table([[p(f'<font color="#168E83"><b>{label.upper()}</b></font>&nbsp;&nbsp; {text}', "body")]], colWidths=[17.9 * cm])
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAF7F5")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#8BD3CB")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return element

    story = []

    # Page 1.
    story += title_band("Page 1 | Executive overview & business layer", "AegisDesk AI", "An evidence-first, human-accountable multi-agent copilot for IT service operations")
    story += [p("Application overview", "h1"), p("AegisDesk serves an ecommerce company's IT, Ecommerce Platform, Data Engineering and BI & Analytics teams. Workplace incidents, checkout/order problems, pipeline failures and Power BI/KPI questions are routed to accountable teams with a visible similar-case solution and approved diagnostic guidance. Five agents operate under deterministic orchestration and human approval."), p("Business problem and value", "h2")]
    story += bullets(["Incomplete tickets create repeated clarification and slow first response.", "Scattered knowledge causes inconsistent troubleshooting and rework.", "Opaque AI introduces hallucination, privacy and automation-bias risk.", "AegisDesk makes evidence and controls visible while reducing repetitive reading, search and drafting."])
    story += [p("Users, stakeholders and process", "h2"), p("Primary users: IT support, Ecommerce Platform, Data Engineering, BI & Analytics and service-desk managers. Beneficiaries: employees and ecommerce operations. Proposed sponsors: Head of IT Support and Ecommerce Technology lead. Control stakeholders include security, privacy/legal, KPI owners and knowledge owners."), callout("Target process", "Ticket intake -> safety -> triage + team route -> article/case search -> visible solution + tips -> draft -> quality validation -> human approval or escalation."), p("Pilot value scorecard", "h2"), table(["Metric", "4-week target", "Evidence"], [["Time to first qualified draft", "-40% vs baseline", "Open-to-approval timestamps"], ["Triage rework", "-25%", "Supervisor correction audit"], ["Approved-evidence coverage", ">=85% covered intents", "Workflow telemetry"], ["Acceptance without material rewrite", ">=60%", "Approval/edit events"], ["Unsafe autonomous sends", "0", "Release-control audit"]], [6.1*cm, 4.1*cm, 7.7*cm]), p("ROI = minutes saved x eligible volume x loaded labor cost - model, platform, integration, training and governance cost. All figures are pilot targets, not measured production results.", "small"), PageBreak()]

    # Page 2.
    story += title_band("Page 2 | Agentic / application layer", "Five roles. One bounded workflow.", "Role separation makes every decision testable, observable and governable")
    story += [p("Agent ecosystem", "h1"), table(["Agent", "Input -> output", "Authority / boundary"], [["Safety & Privacy", "Raw ticket -> sanitized ticket or stop", "Redact, truncate, stop; no generation"], ["Triage", "Ticket -> summary, gaps, category, priority", "Role-specific LLM/mock; fixed taxonomy"], ["Knowledge", "Ticket + triage -> passages and scores", "Approved articles + verified cases; top-3"], ["Resolution", "Triage + evidence -> draft and citations", "Evidence-only LLM/mock; no send"], ["Risk & Quality", "Decisions + draft -> pass/block/escalate", "Deterministic checks; no self-approval"]], [3.3*cm, 7.4*cm, 7.2*cm]), p("Orchestration and decision logic", "h2"), callout("Visible flow", "INTAKE + PRIORITY -> SAFETY -> TRIAGE -> KNOWLEDGE -> RESOLUTION -> QUALITY -> HUMAN")]
    story += bullets(["STOP: prompt-injection indicators end downstream processing and require security review.", "ESCALATE: no evidence above threshold routes to a specialist.", "AWAIT INFORMATION: three or more material gaps produce clarification questions.", "HUMAN REVIEW: every valid draft waits for approval; high/critical impact has an explicit reason.", "BOUND: maximum six trace steps, no arbitrary tools/web, no autonomous ticket change/send; resolution publishing requires a human."])
    story += [p("Human-in-the-loop", "h2"), p("Page 1 analyses employee requests before submission and supports safe self-service or routed ticket creation. Page 2 gives specialists the queue/SLA, evidence, editable final response, quality, trace, and approve/reject/reassign/escalate controls. Publishing a verified resolution remains a separate human action."), callout("Differentiator", "AegisDesk exposes bounded handoffs, accepted evidence, stop/escalate logic and accountable release instead of hiding all reasoning behind one chat response."), PageBreak()]

    # Page 3.
    story += title_band("Page 3 | LLM & data layers", "Language intelligence grounded in approved evidence", "Use an LLM where language adds value; keep critical controls deterministic")
    story += [p("LLM layer", "h1"), p("Two narrow model roles are supported. The Triage Agent returns structured fields using an allowed taxonomy; deterministic policy maps the result to an accountable team. The Resolution Agent receives triage plus accepted evidence and returns a concise draft. Safety, retrieval, validation, routing and approval remain controlled."), table(["Decision", "Current PoC", "Rationale/control"], [["Provider", "Mock; optional Azure/OpenAI-compatible", "Reproducible default; swappable managed path"], ["Prompting", "Two role-specific strict-JSON prompts", "Separate triage and drafting responsibilities"], ["Parameters", "Temperature <=0.4; <=800 output tokens", "Stable, concise and cost-bounded"], ["Failure", "25 s timeout + mock fallback", "Continue safely and expose fallback"], ["Selection", "Benchmark planned", "Accuracy, latency, privacy and cost"]], [3.2*cm, 6.4*cm, 8.3*cm]), p("Data layer and retrieval", "h2"), p("Local SQLite stores runtime tickets, activity, article metadata and verified resolutions. Transparent fixtures seed 27 tickets, 20 approved articles, 16 verified fixes and 15 governed routing evaluations across company departments. Tickets remain restricted operational data."), callout("RAG flow", "Articles + verified cases -> TF-IDF -> triage query -> cosine score >=0.08 -> top-3 -> visible solved-case fix + documentation + grounded draft. Retrieval does not train the LLM.")]
    story += bullets(["No open-web retrieval and no vector database in the current PoC.", "Subject, description and requester are sanitized before cloud inference.", "Low-scoring evidence is rejected rather than labelled as grounded.", "Pilot: RBAC, encryption, version metadata, retention/deletion, residency and ingestion validation.", "Adopt embeddings/hybrid retrieval only if a labelled benchmark proves material benefit."])
    story += [PageBreak()]

    # Page 4.
    story += title_band("Page 4 | Governance, risk & compliance", "Trust is a workflow, not a slogan", "Technical controls, ownership, evidence and residual-risk decisions")
    story += [p("Risk and control register", "h1"), table(["Risk", "Implemented control", "Residual / pilot action"], [["Prompt injection", "Pre-agent detection and stop", "Adversarial and indirect-injection tests"], ["Sensitive disclosure", "Field PII redaction; no raw logs", "DPIA, DLP, retention/residency, incident process"], ["Hallucination", "Approved KB, threshold, citations, quality check", "Groundedness benchmark + sampled QA"], ["Wrong triage", "Taxonomy, confidence, human review", "Labelled accuracy and drift monitoring"], ["Unsafe output", "Prohibited language + provenance checks", "Expanded classifier/red-team suite"], ["Automation bias", "Visible trace/evidence + approval", "Training and edit/override audit"], ["Provider failure", "Timeout, fallback, visible status", "Circuit breaker and availability SLO"]], [3.3*cm, 7.2*cm, 7.4*cm]), p("Governance operating model", "h2")]
    story += bullets(["Accountable product owner; co-owners in security, privacy/legal, knowledge and platform operations.", "Weekly pilot review; monthly release gate.", "Prompt/model/config changes require version, benchmark, sign-off, threshold and rollback criterion.", "Incidents record impact, owner, containment, correction and prevention.", "Priority follows operational impact, not seniority/sensitive traits; assess channel, language and accessibility disparities."])
    story += [p("Compliance posture", "h2"), p("The design applies privacy-by-design and data-minimization principles, but this classroom PoC claims no GDPR, AI Act, ISO, SOC 2 or other certification. A pilot requires organizational legal assessment, vendor terms, processing records, access/retention decisions and security testing."), callout("Release rule", "Risk & Quality may release a draft only to human consideration; it can never authorize customer delivery."), PageBreak()]

    # Page 5.
    story += title_band("Page 5 | Operations, monitoring & references", "Measure value, quality, safety and cost together", "Observe every handoff without logging the customer's raw content")
    story += [p("Operational observability", "h1"), table(["Dimension", "Signals", "Pilot response"], [["Reliability", "Success, fallback, p95 latency, availability", "Alert if fallback >5% or p95 >10 s"], ["Quality", "Triage, evidence, groundedness, edits", "Review failed/edited samples weekly"], ["Safety", "Refusal, PII, unauthorized release", "Immediate incident for disclosure/release"], ["Cost", "Prompt/completion tokens; cost/workflow", "Optimize at release gate"], ["Business", "Draft time, rework, SLA, acceptance, CSAT", "Continue/stop/scale at week 4"]], [3.1*cm, 8.2*cm, 6.6*cm]), p("Continuous improvement", "h2"), p("Collect privacy-minimized metrics -> review escalations/material edits -> label root cause -> change one versioned component -> run regression/adversarial evaluation -> approve/reject -> monitor pilot -> roll back when thresholds fail."), p("Feasible rollout", "h2")]
    story += bullets(["Week 0: governance sign-off, baseline, labelled set and training.", "Weeks 1-2: shadow mode on one queue.", "Weeks 3-4: assisted mode with mandatory approval; measure KPIs and incidents.", "Gate: stop, refine or scale on evidence. Roadmap: SSO/RBAC, ITSM API, managed monitoring and evaluated hybrid retrieval."])
    story += [p("References and disclosure", "h2")]
    story += bullets(["Course assignment: LLM & Business Applications - Designing an Agentic LLM-Based Application.", "NIST AI RMF / GenAI Profile - https://www.nist.gov/itl/ai-risk-management-framework", "OWASP GenAI Top 10 - https://genai.owasp.org/llm-top-10/", "EU GDPR 2016/679 - https://eur-lex.europa.eu/eli/reg/2016/679/", "scikit-learn TF-IDF - https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html", "Microsoft Power BI guidance - https://learn.microsoft.com/en-us/power-bi/visuals/power-bi-visualization-kpi | Stripe Checkout - https://docs.stripe.com/checkout/fulfillment", "Streamlit - https://docs.streamlit.io/ | Microsoft Foundry - https://learn.microsoft.com/en-us/azure/foundry/openai/latest"])
    story += [callout("AI/tool acknowledgement", "Generative AI assisted code, document structure and wording. The student reviewed/tested all claims. Python, Streamlit, scikit-learn, python-docx, python-pptx and ReportLab were used. Tickets/KB are synthetic."), p(f"Prepared by {AUTHOR} | {PROGRAM} | {DATE_LABEL}", "small")]

    pdf.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def build_architecture_pdf(path: Path) -> None:
    """Build a fixed five-page PDF matching the light DOCX architecture report."""
    base = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle("V6Body", parent=base["BodyText"], fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#17324D"), spaceAfter=5),
        "h1": ParagraphStyle("V6H1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#14324F"), spaceBefore=5, spaceAfter=5),
        "h2": ParagraphStyle("V6H2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#078F91"), spaceBefore=5, spaceAfter=3),
        "small": ParagraphStyle("V6Small", parent=base["BodyText"], fontName="Helvetica", fontSize=7.4, leading=9.5, textColor=colors.HexColor("#526B82"), spaceAfter=2),
        "bullet": ParagraphStyle("V6Bullet", parent=base["BodyText"], fontName="Helvetica", fontSize=8.6, leading=11, leftIndent=12, firstLineIndent=-8, textColor=colors.HexColor("#17324D"), spaceAfter=3),
    }
    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=1.35 * cm, leftMargin=1.35 * cm,
        topMargin=1.15 * cm, bottomMargin=1.15 * cm,
        title="AegisDesk AI - Architecture Report", author=AUTHOR,
        subject="Company-wide agentic request platform - six-layer architecture",
        pageCompression=0,
    )

    def header_footer(canvas, doc):
        canvas.saveState()
        width, height = A4
        canvas.setFillColor(colors.HexColor("#DDF2FC"))
        canvas.rect(0, height - 0.32 * cm, width, 0.32 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#078F91"))
        canvas.rect(0, height - 0.32 * cm, 1.1 * cm, 0.32 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#60778C"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(1.35 * cm, 0.52 * cm, "AEGISDESK AI  |  COMPANY-WIDE REQUEST INTELLIGENCE")
        canvas.drawRightString(width - 1.35 * cm, 0.52 * cm, f"PAGE {doc.page}")
        canvas.restoreState()

    def p(text, style="body"):
        return Paragraph(text, styles[style])

    def bullets(items):
        return [p(f"&#8226;&nbsp; {item}", "bullet") for item in items]

    def banner(kicker, title, subtitle):
        content = Paragraph(
            f'<font color="#078F91" size="8"><b>{kicker.upper()}</b></font><br/>'
            f'<font color="#14324F" size="18"><b>{title}</b></font><br/>'
            f'<font color="#526B82" size="8.5">{subtitle}</font>',
            ParagraphStyle("V6Banner", parent=styles["body"], leading=15, spaceAfter=0),
        )
        element = Table([[content]], colWidths=[17.9 * cm])
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF8FF")),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B3D5E7")),
            ("LEFTPADDING", (0, 0), (-1, -1), 11), ("RIGHTPADDING", (0, 0), (-1, -1), 11),
            ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ]))
        return [element, Spacer(1, 6)]

    def card(label, text, alternate=False):
        content = p(f'<font color="#078F91"><b>{label.upper()}</b></font>&nbsp;&nbsp; {text}')
        element = Table([[content]], colWidths=[17.9 * cm])
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3FAFF" if alternate else "#EAF7F5")),
            ("BOX", (0, 0), (-1, -1), 0.55, colors.HexColor("#A9D4D3" if not alternate else "#B8D7E8")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return [element, Spacer(1, 4)]

    def flow():
        cells = [[p(f"<b>{name}</b>", "small") for name in ["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"]]]
        element = Table(cells, colWidths=[17.9 * cm / 6])
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF8FF")),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#9FC7DC")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8D7E8")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        return element

    story = []
    story += banner("Page 1 | Business layer", "AegisDesk AI", "One intelligent front door for every internal company request")
    story += [p("1. Application Overview", "h1"), p("AegisDesk is a company-wide request intelligence platform. Any department can ask another department for support, clarification, correction, or new work. Employees explain the need in plain language; coordinated AI agents clarify it, assess impact, recommend the accountable team, retrieve approved knowledge and verified resolutions, and prepare the next best action.")]
    story += card("Product promise", "The right context, the right answer and the right team - before delays begin.")
    story += [p("2. Key Challenges Addressed", "h2")]
    story += bullets(["Incomplete requests create repeated clarification loops.", "Wrong routing delays ownership and hides urgent work.", "Answers are spread across documents, cases and people, so work is rediscovered.", "Unsupported AI answers create trust, privacy and safety risks."])
    story += [p("3. Core Capabilities (Services Delivered)", "h2")]
    story += bullets(["Employee portal for clarification and safe self-service.", "Cross-department classification, priority and accountable routing.", "Approved knowledge and verified-case search with cited response drafting.", "Specialist workspace with SLA, evidence, editing, approval, reassignment and escalation."])
    story += [p("4. Differentiation", "h2"), p("Five role-separated agents, deterministic orchestration, a real LLM connection and governed RAG create visible intelligence rather than a black-box chatbot. The design integrates with company systems, scales across departments and keeps consequential actions under human control."), PageBreak()]

    story += banner("Page 2 | Agentic / Application layer", "Coordinated intelligence with clear responsibility", "Five specialist agents collaborate; the orchestrator controls every handoff")
    story += [p("Controlled workflow", "h1"), flow(), Spacer(1, 6), p("Five separated responsibilities", "h2")]
    story += card("Safety & Privacy Agent", "Screens input, minimises personal data and stops prompt attacks before cloud inference.")
    story += card("Triage Agent", "Uses the live LLM to summarise, find missing facts and classify the request; approved policy maps it to a team.", True)
    story += card("Knowledge Agent", "Searches approved documentation and human-verified resolutions; weak matches are rejected.")
    story += card("Resolution Agent", "Uses the live LLM to draft from accepted evidence only; it cannot send or execute actions.", True)
    story += card("Risk & Quality Agent", "Checks taxonomy, evidence provenance, unsafe language and quality before human review.")
    story += [p("Controlled autonomy", "h2")]
    story += bullets(["Stop unsafe input; ask for missing facts; escalate weak evidence.", "Route through the approved cross-department taxonomy and prioritise by impact.", "Release only to a human decision point; no arbitrary tools, open web, autonomous send or self-approval."])
    story += [p("A team can be a requester in one workflow and the responsible resolver in another.", "body"), PageBreak()]

    story += banner("Page 3 | LLM and Data layers", "Real language intelligence, grounded in company evidence", "The LLM interprets and drafts; policy, retrieval and release remain controlled")
    story += [p("LLM layer", "h1"), p("AegisDesk defaults to a real OpenAI-compatible deployment. The OpenAI path calls the Responses API with strict JSON Schema outputs, storage disabled, a bounded output budget and a timeout. Only Triage and Resolution are model-backed. Missing credentials or provider failure are visible; live mode never silently substitutes a simulated result.")]
    story += bullets(["Triage returns summary, missing facts, approved category, priority and confidence.", "Resolution uses only accepted internal evidence and asks or escalates instead of inventing.", "Routing, thresholds, validation, approval and audit remain application controls.", "Identifiers are redacted before provider calls; raw content is excluded from telemetry."])
    story += card("Live configuration", "OpenAI Responses API | configured model: gpt-5.6-terra | secret supplied only through OPENAI_API_KEY")
    story += [p("Data and RAG layer", "h1"), p("In operation, authorised sources can include Confluence pages, service records and human-verified resolutions. Retrieval filters by permission, approval, ownership and freshness; only accepted excerpts and source metadata reach the Resolution Agent.")]
    story += card("RAG sequence", "Request and triage context -> permission-aware search -> relevance threshold -> top evidence -> cited draft -> quality validation. Retrieval supplies context; it does not retrain the LLM.", True)
    story += bullets(["Prototype: local SQLite and transparent TF-IDF/cosine retrieval over a demonstration corpus.", "Production: enterprise database, SSO/RBAC, encryption and benchmarked hybrid/vector retrieval."])
    story += [PageBreak()]

    story += banner("Page 4 | Governance, Risk and Compliance layer", "Trust is built into the workflow", "Visible evidence, bounded authority, accountable owners and measurable controls")
    story += [p("Risk-to-control design", "h1")]
    story += card("Incorrect or invented answer", "Approved evidence, visible sources, thresholding, structured output, independent validation and human approval.")
    story += card("Sensitive information", "PII redaction, secret isolation, permission-aware retrieval and no raw-content telemetry.", True)
    story += card("Prompt or document injection", "Pre-agent screening, no arbitrary tools, untrusted attachment handling and adversarial evaluation.")
    story += card("Wrong team or priority", "Fixed taxonomy, policy routing, impact scoring, visible confidence and easy reassignment.", True)
    story += card("Automation bias", "Evidence and uncertainty are visible; specialists can edit, reject or escalate and AI cannot approve itself.")
    story += [p("Governance operating model", "h2")]
    story += bullets(["Product owner owns outcomes; knowledge owners own accuracy and review dates.", "Security, privacy/legal and platform owners approve provider, access, retention, residency and incident controls.", "Prompt, model, taxonomy and threshold changes are versioned, evaluated, approved and reversible.", "Priority uses operational impact, not seniority or sensitive traits."])
    story += [p("Only an authorised person may release a consequential response, reassign ownership, close a case or publish reusable knowledge."), PageBreak()]

    story += banner("Page 5 | Operations, Monitoring and References", "Operate for measurable value", "Quality, reliability, cost and safety are monitored together")
    story += [p("Business and service measures", "h1")]
    story += bullets(["Request completeness, first-time routing accuracy and urgent-case response.", "Time to first useful response, resolution time and SLA compliance.", "Self-service resolution, answer acceptance, specialist edits and knowledge reuse.", "Employee satisfaction and reduction in repeated investigation."])
    story += [p("LLM and operational monitoring", "h2")]
    story += bullets(["Availability, latency, timeout/failure rate, token usage and cost.", "Retrieval relevance, grounding quality, safety stops, PII events and unauthorised-release incidents.", "Review outcomes -> label cause -> version one change -> evaluate -> approve -> monitor or roll back."])
    story += [p("Integration and scale", "h2"), p("Production can connect to Confluence or another knowledge platform, an existing case system, identity/SSO and collaboration channels. Permission metadata travels with indexed content. Stateless services, background ingestion, managed storage and queue-based processing support growth across departments."), p("References and implementation disclosure", "h2")]
    story += bullets(["Course assignment: LLM & Business Applications - Designing an Agentic LLM-Based Application.", "OpenAI Responses API - https://developers.openai.com/api/reference/python/resources/responses/methods/create", "Atlassian Confluence REST API - https://developer.atlassian.com/cloud/confluence/rest/v1/", "NIST AI RMF - https://www.nist.gov/itl/ai-risk-management-framework", "OWASP GenAI Security Project - https://genai.owasp.org/llm-top-10/", "Streamlit - https://docs.streamlit.io/ | scikit-learn - https://scikit-learn.org/"])
    story += card("Disclosure", "Generative AI assisted implementation and drafting; the student reviewed the content and remains responsible. The repository includes local demonstration data. Production requires authorised company sources, credentials, security review and measured results.")
    story += [p(f"Prepared by {AUTHOR} | {PROGRAM} | {DATE_LABEL}", "small")]
    pdf.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


LAYER_COLORS = {
    "business": "2452A4",
    "agentic": "7137A6",
    "llm": "078F91",
    "data": "238BC1",
    "governance": "E78A12",
    "operations": "4B941E",
}

LAYER_PALES = {
    "business": "EEF4FF",
    "agentic": "F5EFFF",
    "llm": "EAF8F6",
    "data": "ECF8FE",
    "governance": "FFF6E8",
    "operations": "F0F8EA",
}


def build_catalog_architecture_document(path: Path) -> None:
    """Build a five-page landscape catalogue with one visual identity per layer."""
    doc = Document()
    configure_doc(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(0.9)
    section.right_margin = Cm(0.9)
    doc.styles["Normal"].font.size = Pt(9)

    def header(page: str, title: str, subtitle: str, color: str) -> None:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width = Cm(2.1)
        table.columns[1].width = Cm(25.4)
        shade_cell(table.cell(0, 0), color)
        shade_cell(table.cell(0, 1), "F2F9FD")
        set_cell_text(table.cell(0, 0), page, bold=True, color=DOC_WHITE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell = table.cell(0, 1)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(title)
        r1.bold = True; r1.font.name = "Aptos Display"; r1.font.size = Pt(20); r1.font.color.rgb = DocxRGB.from_string(DOC_NAVY)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(subtitle)
        r2.font.name = "Aptos"; r2.font.size = Pt(9); r2.font.color.rgb = DOC_MUTED
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def band(number: str, name: str, question: str, color: str) -> None:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.autofit = False
        widths = [Cm(1.1), Cm(6.1), Cm(20.3)]
        for idx, width in enumerate(widths):
            table.columns[idx].width = width
            shade_cell(table.cell(0, idx), color if idx < 2 else "F7FBFE")
        set_cell_text(table.cell(0, 0), number, bold=True, color=DOC_WHITE, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(0, 1), name.upper(), bold=True, color=DOC_WHITE, size=9)
        set_cell_text(table.cell(0, 2), question, bold=True, color=color, size=9)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def cards(items: list[tuple[str, str]], color: str, pale: str) -> None:
        table = doc.add_table(rows=1, cols=len(items))
        table.style = "Table Grid"
        table.autofit = False
        for idx, (title, body) in enumerate(items):
            cell = table.cell(0, idx)
            shade_cell(cell, pale)
            cell.text = ""
            p1 = cell.paragraphs[0]
            p1.paragraph_format.space_after = Pt(3)
            r1 = p1.add_run(title.upper())
            r1.bold = True; r1.font.name = "Aptos"; r1.font.size = Pt(8); r1.font.color.rgb = DocxRGB.from_string(color)
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(body)
            r2.font.name = "Aptos"; r2.font.size = Pt(8.4); r2.font.color.rgb = DocxRGB.from_string(DOC_NAVY)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def flow(items: list[str], color: str, pale: str) -> None:
        table = doc.add_table(rows=1, cols=len(items))
        table.style = "Table Grid"
        table.autofit = False
        for idx, text_value in enumerate(items):
            shade_cell(table.cell(0, idx), color if idx in {0, len(items) - 1} else pale)
            set_cell_text(
                table.cell(0, idx), text_value, bold=True,
                color=DOC_WHITE if idx in {0, len(items) - 1} else color,
                size=7.7, align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def note(title: str, body: str, color: str, pale: str) -> None:
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        shade_cell(table.cell(0, 0), pale)
        cell = table.cell(0, 0)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{title.upper()}  ")
        r.bold = True; r.font.color.rgb = DocxRGB.from_string(color)
        p.add_run(body)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Page 1 — catalogue overview.
    header("CATALOGUE", "AegisDesk AI Architecture", "A simple visual guide to how business value, agents, the LLM, data and control work together", DOC_NAVY)
    note("Application overview", "AegisDesk is the company-wide front door for internal requests. Any department can ask another department for help, clarification, correction or new work. The platform improves the request, finds useful company knowledge and sends it to the right team.", DOC_MINT, "EAF8F6")
    cards([
        ("Requesters", "Employees and teams from every department."),
        ("Resolver teams", "IT, Data, Operations, Finance, HR, Security, Ecommerce and specialists."),
        ("Business value", "Faster answers, fewer transfers, complete requests and reusable knowledge."),
    ], DOC_BLUE, "EEF6FC")
    add_heading(doc, "The experience in five clear steps", 2)
    flow(["1  DESCRIBE", "2  CLARIFY", "3  FIND KNOWLEDGE", "4  ROUTE", "5  RESOLVE"], DOC_MINT, "EAF8F6")
    add_heading(doc, "Architecture catalogue", 2)
    catalogue = doc.add_table(rows=2, cols=3)
    catalogue.style = "Table Grid"
    layers = [
        ("1  BUSINESS", "Why the platform matters", "business"),
        ("2  AGENTIC APPLICATION", "How specialist agents collaborate", "agentic"),
        ("3  LLM", "Where GPT-5.6 Terra adds intelligence", "llm"),
        ("4  DATA", "Where tickets and company knowledge live", "data"),
        ("5  GOVERNANCE & RISK", "How trust and accountability are protected", "governance"),
        ("6  OPERATIONS & MONITORING", "How value and quality improve", "operations"),
    ]
    for idx, (title, body, key) in enumerate(layers):
        cell = catalogue.cell(idx // 3, idx % 3)
        shade_cell(cell, LAYER_PALES[key])
        cell.text = ""
        p1 = cell.paragraphs[0]; p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(title); r1.bold = True; r1.font.size = Pt(8.5); r1.font.color.rgb = DocxRGB.from_string(LAYER_COLORS[key])
        p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(body); r2.font.size = Pt(8.2); r2.font.color.rgb = DocxRGB.from_string(DOC_NAVY)
    note("Design principle", "Use the LLM for language understanding. Use approved rules for routing and control. Keep people accountable for final decisions.", DOC_NAVY, "F1F6FA")
    doc.add_page_break()

    # Page 2 — Business and Agentic layers as two distinct catalogue bands.
    header("PAGE 2", "From business need to coordinated action", "Two clearly separated layers connect company value to agent collaboration", LAYER_COLORS["business"])
    band("1", "Business layer", "Why does AegisDesk exist?", LAYER_COLORS["business"])
    cards([
        ("Problem", "Incomplete requests, wrong routing and scattered answers slow every department."),
        ("Value", "A clearer request reaches the right team with useful context from the start."),
        ("Outcome", "Faster response, less repeated work, better SLA performance and stronger knowledge reuse."),
        ("Measure", "Completeness, first-time routing, response time, resolution time and employee satisfaction."),
    ], LAYER_COLORS["business"], LAYER_PALES["business"])
    flow(["EMPLOYEE OR TEAM", "CLEAR REQUEST", "RIGHT OWNER", "USEFUL ANSWER", "VERIFIED RESOLUTION"], LAYER_COLORS["business"], LAYER_PALES["business"])
    band("2", "Agentic / Application layer", "How do the agents work together?", LAYER_COLORS["agentic"])
    flow(["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"], LAYER_COLORS["agentic"], LAYER_PALES["agentic"])
    cards([
        ("Safety", "Protects sensitive information and stops unsafe instructions."),
        ("Triage", "Clarifies the need, priority and destination team."),
        ("Knowledge", "Finds approved documents and similar verified cases."),
        ("Resolution", "Prepares a useful answer from the evidence."),
        ("Quality", "Checks the answer before a person decides."),
    ], LAYER_COLORS["agentic"], LAYER_PALES["agentic"])
    note("Controlled autonomy", "Agents analyse and recommend. People approve, edit, reassign, escalate and close. A department may be a requester in one case and the resolver in another.", LAYER_COLORS["agentic"], LAYER_PALES["agentic"])
    doc.add_page_break()

    # Page 3 — LLM layer only.
    header("PAGE 3", "LLM Layer", "The language intelligence used by AegisDesk", LAYER_COLORS["llm"])
    band("3", "LLM layer", "How does the application understand and write?", LAYER_COLORS["llm"])
    note("LLM used", "AegisDesk uses GPT-5.6 Terra through the OpenAI Responses API.", LAYER_COLORS["llm"], LAYER_PALES["llm"])
    cards([
        ("Understands", "Reads a request written in normal business language and identifies its meaning."),
        ("Clarifies", "Detects missing information and proposes simple questions for the requester."),
        ("Classifies", "Suggests the request type, urgency and responsible department."),
        ("Drafts", "Prepares an answer using the company evidence selected by the Knowledge Agent."),
    ], LAYER_COLORS["llm"], LAYER_PALES["llm"])
    add_heading(doc, "Two focused uses of the LLM", 2)
    flow(["SAFE REQUEST", "GPT-5.6 TERRA · TRIAGE", "SUMMARY + MISSING FACTS", "CATEGORY + PRIORITY", "TEAM RECOMMENDATION"], LAYER_COLORS["llm"], LAYER_PALES["llm"])
    flow(["APPROVED EVIDENCE", "GPT-5.6 TERRA · RESOLUTION", "DRAFT ANSWER", "QUALITY CHECK", "HUMAN DECISION"], LAYER_COLORS["llm"], LAYER_PALES["llm"])
    cards([
        ("What the LLM decides", "Language meaning, a proposed classification and an evidence-based draft."),
        ("What the LLM cannot do", "It cannot approve itself, send a final answer, change access or close a request."),
        ("Why this creates value", "Teams spend less time rewriting, searching and asking the same clarification questions."),
    ], LAYER_COLORS["llm"], "F6FBFA")
    note("Live operation", "A real OpenAI key is required. If the model is unavailable, AegisDesk shows the failure and does not pretend that a simulated answer came from the LLM.", LAYER_COLORS["llm"], LAYER_PALES["llm"])
    doc.add_page_break()

    # Page 4 — Data layer only, with explicit ticket database.
    header("PAGE 4", "Data Layer", "Where requests, knowledge and verified experience are stored", LAYER_COLORS["data"])
    band("4", "Data layer", "What data is used and where is it stored?", LAYER_COLORS["data"])
    cards([
        ("Ticket database", "SQLite stores every local ticket: description, requester, status, priority, owner, SLA, activity and verified resolution."),
        ("Company knowledge", "Approved procedures and guidance are indexed so the Knowledge Agent can find the most relevant passages."),
        ("Resolved-case memory", "Only solutions confirmed by a person become reusable examples for future requests."),
    ], LAYER_COLORS["data"], LAYER_PALES["data"])
    add_heading(doc, "How data becomes a useful answer", 2)
    flow(["NEW REQUEST", "SEARCH COMPANY KNOWLEDGE", "FIND SIMILAR VERIFIED CASE", "SELECT RELEVANT EVIDENCE", "GROUNDED ANSWER"], LAYER_COLORS["data"], LAYER_PALES["data"])
    cards([
        ("RAG in simple words", "AegisDesk finds the right company information first, then gives it to the LLM to prepare the answer."),
        ("Not model training", "New tickets do not automatically retrain GPT. Verified solutions improve the searchable knowledge available at answer time."),
        ("Access control", "In production, employees see only documents and cases they are authorised to access."),
    ], LAYER_COLORS["data"], "F7FBFE")
    add_heading(doc, "From prototype to company operation", 2)
    flow(["TODAY · SQLITE", "COMPANY PILOT", "SSO + PERMISSIONS", "MANAGED DATABASE", "CONFLUENCE / COMPANY SOURCES"], LAYER_COLORS["data"], LAYER_PALES["data"])
    note("Plain-language summary", "The database remembers the work. The knowledge base explains company rules. RAG finds the best evidence. The LLM turns that evidence into a clear draft.", LAYER_COLORS["data"], LAYER_PALES["data"])
    doc.add_page_break()

    # Page 5 — Governance and Operations, visually separate, with references last.
    header("PAGE 5", "Trust, Operations and Continuous Value", "The final two layers keep AegisDesk safe, reliable and useful", LAYER_COLORS["governance"])
    band("5", "Governance, Risk & Compliance", "How are risk and trust managed?", LAYER_COLORS["governance"])
    cards([
        ("Privacy", "Sensitive details are reduced before the LLM and raw ticket text is excluded from monitoring logs."),
        ("Accuracy", "Answers show their evidence; weak matches are escalated instead of presented as facts."),
        ("Security", "Unsafe instructions are stopped and company permissions protect restricted knowledge."),
        ("Accountability", "A person approves important actions and every decision is traceable."),
    ], LAYER_COLORS["governance"], LAYER_PALES["governance"])
    band("6", "Operations & Monitoring", "How do we run and improve the application?", LAYER_COLORS["operations"])
    cards([
        ("Business value", "Routing accuracy, response time, resolution time, self-service and satisfaction."),
        ("LLM quality", "Useful answers, relevant sources, human edits and model failures."),
        ("Service health", "Availability, speed, safety events and operating cost."),
    ], LAYER_COLORS["operations"], LAYER_PALES["operations"])
    flow(["COLLECT RESULTS", "REVIEW QUALITY", "IMPROVE ONE PART", "TEST", "RELEASE OR ROLLBACK"], LAYER_COLORS["operations"], LAYER_PALES["operations"])
    note("References", "Course assignment · OpenAI Responses API · Atlassian Confluence REST API · NIST AI Risk Management Framework · OWASP GenAI Security Project · Streamlit · scikit-learn.", DOC_NAVY, "F1F6FA")
    note("Disclosure", "Generative AI assisted implementation and drafting; the student reviewed the work and remains responsible. The repository uses local demonstration data. Company operation requires authorised sources, credentials, permissions, security review and measured results.", DOC_NAVY, "F1F6FA")
    add_body(doc, f"Prepared by {AUTHOR} · {PROGRAM} · {DATE_LABEL}")
    doc.save(path)


def build_catalog_architecture_pdf(path: Path) -> None:
    """Build the matching fixed five-page landscape catalogue PDF."""
    page_size = rl_landscape(A4)
    base = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle("CatBody", parent=base["BodyText"], fontName="Helvetica", fontSize=8, leading=10.2, textColor=colors.HexColor("#17324D"), spaceAfter=3),
        "title": ParagraphStyle("CatTitle", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=19, leading=22, textColor=colors.HexColor("#14324F"), spaceAfter=2),
        "h2": ParagraphStyle("CatH2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=colors.HexColor("#14324F"), spaceBefore=3, spaceAfter=3),
        "small": ParagraphStyle("CatSmall", parent=base["BodyText"], fontName="Helvetica", fontSize=6.8, leading=8.3, textColor=colors.HexColor("#526B82"), spaceAfter=1),
    }
    pdf = SimpleDocTemplate(
        str(path), pagesize=page_size, leftMargin=0.85 * cm, rightMargin=0.85 * cm,
        topMargin=0.7 * cm, bottomMargin=0.7 * cm,
        title="AegisDesk AI - Visual Architecture Catalogue", author=AUTHOR,
        subject="Five-page visual catalogue of the six-layer architecture", pageCompression=0,
    )
    page_width, page_height = page_size
    content_width = page_width - 1.7 * cm

    def p(text, style="body"):
        return Paragraph(text, styles[style])

    def header(page_label, title, subtitle, color):
        page_cell = p(f'<font color="#FFFFFF" size="9"><b>{page_label}</b></font>')
        title_cell = Paragraph(
            f'<font color="#14324F" size="19"><b>{title}</b></font><br/>'
            f'<font color="#526B82" size="8">{subtitle}</font>',
            ParagraphStyle("CatHeader", parent=styles["body"], leading=15, spaceAfter=0),
        )
        table = Table([[page_cell, title_cell]], colWidths=[2.0 * cm, content_width - 2.0 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(f"#{color}")),
            ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#F2F9FD")),
            ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#B8D5E5")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, 0), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        return [table, Spacer(1, 5)]

    def band(number, name, question, color):
        table = Table(
            [[p(f'<font color="#FFFFFF"><b>{number}</b></font>'), p(f'<font color="#FFFFFF"><b>{name.upper()}</b></font>'), p(f'<font color="#{color}"><b>{question}</b></font>')]],
            colWidths=[0.9 * cm, 5.7 * cm, content_width - 6.6 * cm],
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (1, 0), colors.HexColor(f"#{color}")),
            ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#F7FBFE")),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(f"#{color}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return [table, Spacer(1, 4)]

    def cards(items, color, pale):
        table = Table(
            [[Paragraph(f'<font color="#{color}" size="7"><b>{title.upper()}</b></font><br/><font color="#17324D" size="7.6">{body}</font>', styles["small"]) for title, body in items]],
            colWidths=[content_width / len(items)] * len(items),
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{pale}")),
            ("BOX", (0, 0), (-1, -1), 0.55, colors.HexColor(f"#{color}")),
            ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{color}")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return [table, Spacer(1, 4)]

    def flow(items, color, pale):
        content = []
        for idx, item in enumerate(items):
            fg = "#FFFFFF" if idx in {0, len(items) - 1} else f"#{color}"
            content.append(p(f'<font color="{fg}" size="7"><b>{item}</b></font>', "small"))
        table = Table([content], colWidths=[content_width / len(items)] * len(items))
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(f"#{color}")),
            ("BACKGROUND", (-1, 0), (-1, 0), colors.HexColor(f"#{color}")),
            ("BACKGROUND", (1, 0), (-2, 0), colors.HexColor(f"#{pale}")),
            ("BOX", (0, 0), (-1, -1), 0.55, colors.HexColor(f"#{color}")),
            ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{color}")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return [table, Spacer(1, 4)]

    def note(label, text_value, color, pale):
        element = Table(
            [[p(f'<font color="#{color}"><b>{label.upper()}</b></font>&nbsp;&nbsp; {text_value}')]],
            colWidths=[content_width],
        )
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{pale}")),
            ("BOX", (0, 0), (-1, -1), 0.55, colors.HexColor(f"#{color}")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return [element, Spacer(1, 4)]

    def page_footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D6E4ED"))
        canvas.line(0.85 * cm, 0.52 * cm, page_width - 0.85 * cm, 0.52 * cm)
        canvas.setFillColor(colors.HexColor("#60778C")); canvas.setFont("Helvetica", 6.5)
        canvas.drawString(0.85 * cm, 0.25 * cm, "AEGISDESK AI  |  VISUAL ARCHITECTURE CATALOGUE")
        canvas.drawRightString(page_width - 0.85 * cm, 0.25 * cm, f"PAGE {doc.page} OF 5")
        canvas.restoreState()

    story = []
    story += header("CATALOGUE", "AegisDesk AI Architecture", "A simple visual guide to how business value, agents, the LLM, data and control work together", DOC_NAVY)
    story += note("Application overview", "AegisDesk is the company-wide front door for internal requests. Any department can ask another department for help, clarification, correction or new work. The platform improves the request, finds useful company knowledge and sends it to the right team.", DOC_MINT, "EAF8F6")
    story += cards([("Requesters", "Employees and teams from every department."), ("Resolver teams", "IT, Data, Operations, Finance, HR, Security, Ecommerce and specialists."), ("Business value", "Faster answers, fewer transfers, complete requests and reusable knowledge.")], DOC_BLUE, "EEF6FC")
    story += [p("The experience in five clear steps", "h2")]
    story += flow(["1  DESCRIBE", "2  CLARIFY", "3  FIND KNOWLEDGE", "4  ROUTE", "5  RESOLVE"], DOC_MINT, "EAF8F6")
    story += [p("Architecture catalogue", "h2")]
    layer_cells = []
    for title, body, key in [("1 BUSINESS", "Why the platform matters", "business"), ("2 AGENTIC APPLICATION", "How specialist agents collaborate", "agentic"), ("3 LLM", "Where GPT-5.6 Terra adds intelligence", "llm"), ("4 DATA", "Where tickets and knowledge live", "data"), ("5 GOVERNANCE & RISK", "How trust is protected", "governance"), ("6 OPERATIONS & MONITORING", "How value and quality improve", "operations")]:
        layer_cells.append(Paragraph(f'<font color="#{LAYER_COLORS[key]}" size="7.5"><b>{title}</b></font><br/><font color="#17324D" size="7.3">{body}</font>', styles["small"]))
    layer_table = Table([layer_cells[:3], layer_cells[3:]], colWidths=[content_width / 3] * 3)
    layer_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8D5E5")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, -1), colors.white), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
    story += [layer_table, Spacer(1, 4)]
    story += note("Design principle", "Use the LLM for language understanding. Use approved rules for routing and control. Keep people accountable for final decisions.", DOC_NAVY, "F1F6FA")
    story += [PageBreak()]

    story += header("PAGE 2", "From business need to coordinated action", "Two clearly separated layers connect company value to agent collaboration", LAYER_COLORS["business"])
    story += band("1", "Business layer", "Why does AegisDesk exist?", LAYER_COLORS["business"])
    story += cards([("Problem", "Incomplete requests, wrong routing and scattered answers slow every department."), ("Value", "A clearer request reaches the right team with useful context from the start."), ("Outcome", "Faster response, less repeated work and stronger knowledge reuse."), ("Measure", "Completeness, routing, response time, resolution time and satisfaction.")], LAYER_COLORS["business"], LAYER_PALES["business"])
    story += flow(["EMPLOYEE OR TEAM", "CLEAR REQUEST", "RIGHT OWNER", "USEFUL ANSWER", "VERIFIED RESOLUTION"], LAYER_COLORS["business"], LAYER_PALES["business"])
    story += band("2", "Agentic / Application layer", "How do the agents work together?", LAYER_COLORS["agentic"])
    story += flow(["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"], LAYER_COLORS["agentic"], LAYER_PALES["agentic"])
    story += cards([("Safety", "Protects information and stops unsafe instructions."), ("Triage", "Clarifies the need, priority and destination."), ("Knowledge", "Finds approved documents and verified cases."), ("Resolution", "Prepares an answer from the evidence."), ("Quality", "Checks it before a person decides.")], LAYER_COLORS["agentic"], LAYER_PALES["agentic"])
    story += note("Controlled autonomy", "Agents analyse and recommend. People approve, edit, reassign, escalate and close. A department may request in one case and resolve in another.", LAYER_COLORS["agentic"], LAYER_PALES["agentic"])
    story += [PageBreak()]

    story += header("PAGE 3", "LLM Layer", "The language intelligence used by AegisDesk", LAYER_COLORS["llm"])
    story += band("3", "LLM layer", "How does the application understand and write?", LAYER_COLORS["llm"])
    story += note("LLM used", "AegisDesk uses GPT-5.6 Terra through the OpenAI Responses API.", LAYER_COLORS["llm"], LAYER_PALES["llm"])
    story += cards([("Understands", "Reads normal business language and identifies its meaning."), ("Clarifies", "Detects missing information and proposes simple questions."), ("Classifies", "Suggests request type, urgency and department."), ("Drafts", "Prepares an answer from selected company evidence.")], LAYER_COLORS["llm"], LAYER_PALES["llm"])
    story += [p("Two focused uses of the LLM", "h2")]
    story += flow(["SAFE REQUEST", "GPT-5.6 TERRA · TRIAGE", "SUMMARY + MISSING FACTS", "CATEGORY + PRIORITY", "TEAM"], LAYER_COLORS["llm"], LAYER_PALES["llm"])
    story += flow(["APPROVED EVIDENCE", "GPT-5.6 TERRA · RESOLUTION", "DRAFT ANSWER", "QUALITY CHECK", "HUMAN DECISION"], LAYER_COLORS["llm"], LAYER_PALES["llm"])
    story += cards([("What it decides", "Meaning, proposed classification and an evidence-based draft."), ("What it cannot do", "Approve itself, send a final answer, change access or close a request."), ("Business value", "Less rewriting, searching and repeated clarification.")], LAYER_COLORS["llm"], "F6FBFA")
    story += note("Live operation", "A real OpenAI key is required. If the model is unavailable, AegisDesk shows the failure and never presents a simulated answer as live LLM output.", LAYER_COLORS["llm"], LAYER_PALES["llm"])
    story += [PageBreak()]

    story += header("PAGE 4", "Data Layer", "Where requests, knowledge and verified experience are stored", LAYER_COLORS["data"])
    story += band("4", "Data layer", "What data is used and where is it stored?", LAYER_COLORS["data"])
    story += cards([("Ticket database", "SQLite stores each local ticket: description, requester, status, priority, owner, SLA, activity and verified resolution."), ("Company knowledge", "Approved procedures and guidance are indexed for relevant search."), ("Resolved-case memory", "Only solutions confirmed by a person become reusable examples.")], LAYER_COLORS["data"], LAYER_PALES["data"])
    story += [p("How data becomes a useful answer", "h2")]
    story += flow(["NEW REQUEST", "SEARCH KNOWLEDGE", "SIMILAR VERIFIED CASE", "RELEVANT EVIDENCE", "GROUNDED ANSWER"], LAYER_COLORS["data"], LAYER_PALES["data"])
    story += cards([("RAG in simple words", "Find the right company information first, then give it to the LLM."), ("Not model training", "New tickets do not automatically retrain GPT. Verified solutions improve searchable knowledge."), ("Access control", "In production, people see only sources they are authorised to access.")], LAYER_COLORS["data"], "F7FBFE")
    story += [p("From prototype to company operation", "h2")]
    story += flow(["TODAY · SQLITE", "COMPANY PILOT", "SSO + PERMISSIONS", "MANAGED DATABASE", "CONFLUENCE / COMPANY SOURCES"], LAYER_COLORS["data"], LAYER_PALES["data"])
    story += note("Plain-language summary", "The database remembers the work. The knowledge base explains company rules. RAG finds the best evidence. The LLM turns it into a clear draft.", LAYER_COLORS["data"], LAYER_PALES["data"])
    story += [PageBreak()]

    story += header("PAGE 5", "Trust, Operations and Continuous Value", "The final two layers keep AegisDesk safe, reliable and useful", LAYER_COLORS["governance"])
    story += band("5", "Governance, Risk & Compliance", "How are risk and trust managed?", LAYER_COLORS["governance"])
    story += cards([("Privacy", "Reduce sensitive details before the LLM; exclude raw text from monitoring logs."), ("Accuracy", "Show evidence and escalate weak matches."), ("Security", "Stop unsafe instructions and respect permissions."), ("Accountability", "A person approves important actions; decisions are traceable.")], LAYER_COLORS["governance"], LAYER_PALES["governance"])
    story += band("6", "Operations & Monitoring", "How do we run and improve the application?", LAYER_COLORS["operations"])
    story += cards([("Business value", "Routing, response, resolution, self-service and satisfaction."), ("LLM quality", "Useful answers, relevant sources, human edits and failures."), ("Service health", "Availability, speed, safety events and operating cost.")], LAYER_COLORS["operations"], LAYER_PALES["operations"])
    story += flow(["COLLECT RESULTS", "REVIEW QUALITY", "IMPROVE ONE PART", "TEST", "RELEASE OR ROLLBACK"], LAYER_COLORS["operations"], LAYER_PALES["operations"])
    story += note("References", "Course assignment · OpenAI Responses API · Atlassian Confluence REST API · NIST AI RMF · OWASP GenAI · Streamlit · scikit-learn.", DOC_NAVY, "F1F6FA")
    story += note("Disclosure", "Generative AI assisted implementation and drafting; the student reviewed the work. The repository uses local demonstration data. Company operation requires authorised sources, credentials, permissions, security review and measured results.", DOC_NAVY, "F1F6FA")
    story += [p(f"Prepared by {AUTHOR} | {PROGRAM} | {DATE_LABEL}", "small")]

    pdf.build(story, onFirstPage=page_footer, onLaterPages=page_footer)


def build_template_architecture_document(path: Path) -> None:
    """Build the report with the exact horizontal-layer logic of the course template."""
    doc = Document()
    configure_doc(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(0.45)
    section.bottom_margin = Cm(0.55)
    section.left_margin = Cm(0.55)
    section.right_margin = Cm(0.55)
    doc.styles["Normal"].font.size = Pt(8)

    def spacer(points=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(points)
        p.paragraph_format.space_before = Pt(0)

    def rich(cell, heading: str, body: str, color: str = DOC_NAVY, *, size=7.2, center=False):
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run(heading.upper())
        r.bold = True; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.color.rgb = DocxRGB.from_string(color)
        if body:
            p2 = cell.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(body)
            r2.font.name = "Aptos"; r2.font.size = Pt(size); r2.font.color.rgb = DocxRGB.from_string(DOC_NAVY)

    def header(title: str, subtitle: str, page_label: str):
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"; table.autofit = False
        shade_cell(table.cell(0, 0), "0A2F5A"); shade_cell(table.cell(0, 1), "0A2F5A")
        table.columns[0].width = Cm(24.5); table.columns[1].width = Cm(4.0)
        cell = table.cell(0, 0); cell.text = ""
        p1 = cell.paragraphs[0]; p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(title); r1.bold = True; r1.font.name = "Aptos Display"; r1.font.size = Pt(18); r1.font.color.rgb = DocxRGB.from_string(DOC_WHITE)
        p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(subtitle); r2.font.name = "Aptos"; r2.font.size = Pt(8.5); r2.font.color.rgb = DocxRGB.from_string("55C9F3")
        rich(table.cell(0, 1), "AEGISDESK AI", page_label, DOC_WHITE, size=7.5, center=True)
        for paragraph in table.cell(0, 1).paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = DocxRGB.from_string(DOC_WHITE)
        spacer(1)

    def overview(summary: str, users: str):
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"; table.autofit = False
        widths = [Cm(6.1), Cm(14.2), Cm(8.0)]
        for i, width in enumerate(widths):
            table.columns[i].width = width
            shade_cell(table.cell(0, i), "F5FAFE" if i else "EEF6FC")
        rich(table.cell(0, 0), "APPLICATION OVERVIEW", "One intelligent front door for every internal company request.", "0A2F5A", size=7.5)
        rich(table.cell(0, 1), "PURPOSE & VALUE", summary, "0A2F5A", size=7.2)
        rich(table.cell(0, 2), "MAIN USERS / STAKEHOLDERS", users, "0A2F5A", size=7.0)
        spacer(1)

    def add_flow(cell, title: str, items: list[str], color: str, pale: str):
        cell.text = ""
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(6.8); r.font.color.rgb = DocxRGB.from_string(color)
        nested = cell.add_table(rows=1, cols=len(items))
        nested.style = "Table Grid"; nested.autofit = False
        for idx, item in enumerate(items):
            shade_cell(nested.cell(0, idx), pale)
            set_cell_text(nested.cell(0, idx), item, bold=True, color=color, size=6.1, align=WD_ALIGN_PARAGRAPH.CENTER)

    expanded_rows = False

    def layer_row(number: str, name: str, question: str, color: str, pale: str,
                  section_title: str, section_body: str, flow_title: str,
                  flow_items: list[str], notes: str, row_height_cm: float | None = None):
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"; table.autofit = False
        if row_height_cm is None:
            row_height_cm = 3.0 if expanded_rows else 2.15
        if row_height_cm is not None:
            table.rows[0].height = Cm(row_height_cm)
            table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        widths = [Cm(5.2), Cm(8.5), Cm(10.0), Cm(4.6)]
        for idx, width in enumerate(widths):
            table.columns[idx].width = width
        shade_cell(table.cell(0, 0), color)
        shade_cell(table.cell(0, 1), "FFFFFF")
        shade_cell(table.cell(0, 2), pale)
        shade_cell(table.cell(0, 3), "F8FBFD")
        left = table.cell(0, 0); left.text = ""
        p1 = left.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER; p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(number); r1.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = DocxRGB.from_string(DOC_WHITE)
        p2 = left.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(name.upper()); r2.bold = True; r2.font.size = Pt(7.4); r2.font.color.rgb = DocxRGB.from_string(DOC_WHITE)
        p3 = left.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER; p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(question); r3.font.size = Pt(6.3); r3.font.color.rgb = DocxRGB.from_string(DOC_WHITE)
        rich(table.cell(0, 1), section_title, section_body, color, size=6.7)
        add_flow(table.cell(0, 2), flow_title, flow_items, color, pale)
        rich(table.cell(0, 3), "EXAMPLES / NOTES", notes, color, size=6.3)
        spacer(1)

    def footer(principles: str, technologies: str, integrations: str):
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"; table.autofit = False
        for idx, width in enumerate([Cm(9.4), Cm(9.4), Cm(9.5)]):
            table.columns[idx].width = width; shade_cell(table.cell(0, idx), "F2F7FA")
        rich(table.cell(0, 0), "ARCHITECTURE PRINCIPLES", principles, "0A2F5A", size=6.2)
        rich(table.cell(0, 1), "KEY TECHNOLOGIES", technologies, "0A2F5A", size=6.2)
        rich(table.cell(0, 2), "EXTERNAL INTEGRATIONS", integrations, "0A2F5A", size=6.2)

    common_summary = "AegisDesk clarifies incomplete requests, finds approved answers and verified solutions, recommends the right department and keeps people responsible for final decisions."
    common_users = "Requesters: every department\nResolvers: IT, Data, Operations, Finance, HR, Security, Ecommerce\nOwners: service, knowledge, security and privacy teams"

    # Page 1 — exact architecture map logic.
    header("ARCHITECTURE – LLM AGENTIC APPLICATION", "How the AegisDesk layers interact to deliver faster, safer cross-department service", "ARCHITECTURE MAP")
    overview(common_summary, common_users)
    layer_row("1", "Business Layer", "Why does this exist?", LAYER_COLORS["business"], LAYER_PALES["business"], "OBJECTIVES & VALUE", "Complete requests\nCorrect routing\nFaster answers\nReusable knowledge", "BUSINESS PROCESS", ["REQUEST", "CLARIFY", "ROUTE", "RESOLVE"], "Value: less delay\nMeasure: time + routing")
    layer_row("2", "Agentic / Application Layer", "What do the agents do?", LAYER_COLORS["agentic"], LAYER_PALES["agentic"], "AGENTS & ORCHESTRATION", "Five specialist roles\nOrdered handoffs\nHuman approval", "AGENT WORKFLOW", ["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY"], "Stop, ask or escalate\nNo autonomous send")
    layer_row("3", "LLM Layer", "How is intelligence provided?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "LLM UTILISATION", "GPT-5.6 Terra\nUnderstands requests\nAsks missing questions\nDrafts grounded answers", "LLM COMPONENTS", ["REQUEST", "GPT-5.6", "PROMPT", "DRAFT"], "Used for language\nHuman reviews output")
    layer_row("4", "Data Layer", "What data is used?", LAYER_COLORS["data"], LAYER_PALES["data"], "DATA MANAGEMENT", "SQLite ticket database\nApproved company documents\nVerified resolutions", "DATA FLOW", ["SOURCES", "SEARCH", "EVIDENCE", "ANSWER"], "RAG, not training\nPermission-aware")
    layer_row("5", "Governance, Risk & Compliance", "How do we manage risk?", LAYER_COLORS["governance"], LAYER_PALES["governance"], "GOVERNANCE & CONTROLS", "Privacy protection\nEvidence and confidence\nHuman accountability", "KEY CONTROLS", ["PRIVACY", "SECURITY", "QUALITY", "APPROVAL"], "Prevent unsafe answers\nKeep audit trail")
    layer_row("6", "Operations & Monitoring", "How do we run and improve it?", LAYER_COLORS["operations"], LAYER_PALES["operations"], "OPERATIONS & MONITORING", "Business value\nLLM quality\nService health\nContinuous improvement", "IMPROVEMENT LOOP", ["COLLECT", "REVIEW", "IMPROVE", "RELEASE"], "Measure results\nRollback when needed")
    footer("Scalable · Secure · Explainable · Human-centred", "Streamlit · GPT-5.6 Terra · SQLite · RAG", "Confluence · Company case system · SSO · Collaboration tools")
    doc.add_page_break()

    # Page 2 — detailed Business and Agentic sections using the same four-column logic.
    expanded_rows = True
    header("BUSINESS & AGENTIC DESIGN", "How AegisDesk turns an unclear request into coordinated action", "PAGE 2 OF 5")
    overview("The platform serves the whole company. A department can request help in one workflow and resolve another department's request in the next.", common_users)
    layer_row("1", "Business Layer", "Why does this exist?", LAYER_COLORS["business"], LAYER_PALES["business"], "OBJECTIVES & VALUE", "Reduce clarification loops\nPrevent wrong-team transfers\nPrioritise high-impact work\nReuse company knowledge", "PROCESS IMPACTED", ["DESCRIBE", "COMPLETE", "ASSIGN", "ANSWER", "LEARN"], "KPIs:\nCompleteness\nRouting accuracy\nResponse time\nResolution time", row_height_cm=3.0)
    layer_row("1A", "Business Value", "What changes for teams?", LAYER_COLORS["business"], LAYER_PALES["business"], "VALUE BY USER", "Requester: faster guidance\nResolver: complete context\nManager: visible ownership\nCompany: reusable learning", "VALUE CHAIN", ["LESS SEARCH", "FEWER TRANSFERS", "FASTER ACTION", "BETTER SERVICE"], "Investor value:\nScales across departments\nUses existing knowledge", row_height_cm=3.0)
    layer_row("2", "Agentic / Application Layer", "What do the agents do?", LAYER_COLORS["agentic"], LAYER_PALES["agentic"], "AGENT ECOSYSTEM", "Safety protects\nTriage understands\nKnowledge searches\nResolution drafts\nQuality validates", "EXAMPLE WORKFLOW", ["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"], "Orchestrator controls order\nAgents have separate roles", row_height_cm=3.0)
    layer_row("2A", "Controlled Autonomy", "Who makes the final decision?", LAYER_COLORS["agentic"], LAYER_PALES["agentic"], "DECISION RULES", "Stop unsafe input\nAsk for missing facts\nEscalate weak evidence\nRoute to approved teams", "HUMAN-IN-THE-LOOP", ["AI RECOMMENDS", "PERSON REVIEWS", "EDIT / REASSIGN", "APPROVE / ESCALATE"], "No autonomous send\nNo self-approval\nVisible agent trace", row_height_cm=3.0)
    footer("Value first · Clear ownership · Controlled autonomy", "Employee Portal · Support Workspace · Five agents", "Company departments · Knowledge sources · Case systems")
    doc.add_page_break()

    # Page 3 — LLM layer, separate and explicit.
    header("LLM LAYER", "Where GPT-5.6 Terra adds language intelligence—and where it does not", "PAGE 3 OF 5")
    overview("AegisDesk uses GPT-5.6 Terra through the OpenAI Responses API. The LLM reads normal business language and prepares structured recommendations.", "Used by: Triage Agent and Resolution Agent\nReviewed by: Risk & Quality Agent and an authorised person")
    layer_row("3", "LLM Layer", "How is intelligence provided?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "LLM UTILISATION", "Understand the request\nSummarise the need\nFind missing information\nPropose classification\nDraft an answer", "LLM COMPONENTS", ["GPT-5.6 TERRA", "ROLE PROMPT", "COMPANY EVIDENCE", "STRUCTURED RESULT"], "Language tasks only\nEvidence is supplied by RAG")
    layer_row("3A", "Triage Use", "How is the request understood?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "INPUT & RESULT", "Input: safe request\nOutput: summary, questions, category, priority and team recommendation", "TRIAGE FLOW", ["REQUEST", "UNDERSTAND", "CLARIFY", "CLASSIFY", "ROUTE"], "Example value:\nLess manual reading\nFewer clarification messages")
    layer_row("3B", "Resolution Use", "How is the answer prepared?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "INPUT & RESULT", "Input: request + approved evidence\nOutput: clear draft with sources and uncertainty", "RESOLUTION FLOW", ["EVIDENCE", "DRAFT", "CHECK", "HUMAN REVIEW"], "No evidence = escalate\nThe LLM must not invent")
    layer_row("3C", "LLM Boundary", "What can the LLM not do?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "CONTROLLED RESPONSIBILITY", "Cannot send the final answer\nCannot grant access\nCannot close a ticket\nCannot approve itself", "CONTROL FLOW", ["LLM SUGGESTS", "RULES CHECK", "PERSON DECIDES"], "If OpenAI is unavailable:\nshow the failure\ndo not fake a live answer")
    footer("Useful · Grounded · Reviewable · Human-controlled", "GPT-5.6 Terra · OpenAI Responses API", "Approved company evidence supplied through RAG")
    doc.add_page_break()

    # Page 4 — Data layer, separate and explicit.
    header("DATA LAYER", "Where every ticket, document and verified solution lives", "PAGE 4 OF 5")
    overview("The Data Layer gives AegisDesk memory. It stores operational tickets and makes approved company knowledge searchable for the agents.", "Data owners: service teams and knowledge owners\nAccess: only authorised employees and resolver teams")
    layer_row("4", "Data Layer", "What data is used?", LAYER_COLORS["data"], LAYER_PALES["data"], "TICKET DATABASE", "SQLite stores:\nTicket description\nRequester and owner\nStatus and priority\nSLA and activity\nVerified resolution", "TICKET LIFE CYCLE", ["CREATE", "STORE", "ASSIGN", "UPDATE", "RESOLVE"], "Current application database:\ndata/aegisdesk.db")
    layer_row("4A", "Knowledge Sources", "Where do answers come from?", LAYER_COLORS["data"], LAYER_PALES["data"], "APPROVED KNOWLEDGE", "Company procedures\nOperational guidance\nData definitions\nTechnical documentation\nVerified past solutions", "SOURCE FLOW", ["CONFLUENCE / DOCS", "APPROVE", "INDEX", "SEARCH"], "Production sources keep owner, date and permissions")
    layer_row("4B", "RAG", "How is the right evidence found?", LAYER_COLORS["data"], LAYER_PALES["data"], "RAG IN SIMPLE WORDS", "Search first\nSelect relevant company evidence\nGive that evidence to GPT-5.6 Terra\nShow the source with the answer", "RAG FLOW", ["REQUEST", "SEARCH", "BEST EVIDENCE", "LLM DRAFT", "CITATION"], "RAG does not retrain GPT\nIt supplies current context")
    layer_row("4C", "Data Protection", "Who can see what?", LAYER_COLORS["data"], LAYER_PALES["data"], "DATA CONTROLS", "User permissions\nApproved-source status\nSensitive-data reduction\nRetention rules\nHuman verification", "ACCESS FLOW", ["IDENTIFY USER", "CHECK ACCESS", "FILTER SOURCES", "RETURN EVIDENCE"], "Future company operation:\nSSO + managed database\npermission-aware Confluence")
    footer("Reliable · Permission-aware · Current · Reusable", "SQLite today · Managed database in production · RAG", "Confluence · Company documents · Existing tickets")
    doc.add_page_break()

    # Page 5 — Governance and Operations, same template logic; references remain last.
    header("GOVERNANCE, OPERATIONS & MONITORING", "How AegisDesk stays trustworthy and improves over time", "PAGE 5 OF 5")
    overview("AegisDesk combines technical safeguards, clear ownership and continuous measurement. AI recommendations remain visible and reversible.", "Accountable owners: product, service, security, privacy, knowledge and platform operations")
    layer_row("5", "Governance, Risk & Compliance", "How do we manage risk?", LAYER_COLORS["governance"], LAYER_PALES["governance"], "RISKS & CONTROLS", "Wrong answer: require evidence\nSensitive data: reduce and protect\nUnsafe instruction: stop\nWrong routing: allow reassignment\nOver-trust: require human review", "KEY CONTROLS", ["PRIVACY", "SECURITY", "EVIDENCE", "QUALITY", "APPROVAL"], "Every important action has an owner and audit trail")
    layer_row("5A", "Human Accountability", "Who is responsible?", LAYER_COLORS["governance"], LAYER_PALES["governance"], "OWNERSHIP", "Knowledge owners approve sources\nSecurity and privacy approve controls\nService owners measure outcomes\nAuthorised people release answers", "DECISION FLOW", ["AI PROPOSES", "CONTROL CHECK", "HUMAN DECIDES", "AUDIT"], "No claim of certification without company assessment")
    layer_row("6", "Operations & Monitoring", "How do we run the application?", LAYER_COLORS["operations"], LAYER_PALES["operations"], "WHAT WE MONITOR", "Business: time, routing, satisfaction\nLLM: usefulness and failures\nRAG: source relevance\nService: speed and availability\nSafety: blocked events", "OBSERVABILITY LOOP", ["COLLECT", "MONITOR", "ALERT", "IMPROVE"], "Improve only after review and testing")
    layer_row("6A", "Continuous Improvement", "How does performance improve?", LAYER_COLORS["operations"], LAYER_PALES["operations"], "IMPROVEMENT CYCLE", "Review outcomes\nFind the cause\nChange one component\nTest quality and safety\nRelease or rollback", "CONTROLLED CHANGE", ["MEASURE", "LEARN", "TEST", "APPROVE", "RELEASE"], "Scale across departments when value is demonstrated")
    footer("References: course assignment · OpenAI · Atlassian · NIST AI RMF · OWASP GenAI", "Streamlit · GPT-5.6 Terra · SQLite · scikit-learn", "Confluence · SSO · Company case system · Collaboration channels")
    add_body(doc, "Disclosure: Generative AI assisted implementation and drafting. The student reviewed the work and remains responsible. Company deployment requires authorised data, credentials, permissions, security review and measured results.")
    doc.save(path)


def build_template_architecture_pdf(path: Path) -> None:
    """Build the matching five-page PDF using the course template's row logic."""
    page_size = rl_landscape(A4)
    page_width, _ = page_size
    content_width = page_width - 1.1 * cm
    base = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle("TplBody", parent=base["BodyText"], fontName="Helvetica", fontSize=6.4, leading=7.8, textColor=colors.HexColor("#17324D"), spaceAfter=0),
        "tiny": ParagraphStyle("TplTiny", parent=base["BodyText"], fontName="Helvetica", fontSize=5.4, leading=6.6, textColor=colors.HexColor("#17324D"), spaceAfter=0),
        "white": ParagraphStyle("TplWhite", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=6.2, leading=7.4, textColor=colors.white, alignment=1, spaceAfter=0),
    }
    pdf = SimpleDocTemplate(
        str(path), pagesize=page_size, leftMargin=0.55 * cm, rightMargin=0.55 * cm,
        topMargin=0.4 * cm, bottomMargin=0.45 * cm,
        title="AegisDesk AI - LLM Agentic Application Architecture",
        author=AUTHOR, subject="Architecture report based on the course layer template", pageCompression=0,
    )

    def p(text, style="body"):
        return Paragraph(text, styles[style])

    def header(title, subtitle, page_label):
        left = Paragraph(f'<font color="#FFFFFF" size="17"><b>{title}</b></font><br/><font color="#55C9F3" size="8">{subtitle}</font>', ParagraphStyle("TplHeader", parent=styles["body"], leading=14))
        right = p(f"AEGISDESK AI<br/>{page_label}", "white")
        table = Table([[left, right]], colWidths=[content_width - 4.0 * cm, 4.0 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0A2F5A")),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#0A2F5A")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return [table, Spacer(1, 2)]

    def overview(summary, users):
        cells = [
            p('<font color="#0A2F5A"><b>APPLICATION OVERVIEW</b></font><br/>One intelligent front door for every internal company request.'),
            p(f'<font color="#0A2F5A"><b>PURPOSE &amp; VALUE</b></font><br/>{summary}'),
            p(f'<font color="#0A2F5A"><b>MAIN USERS / STAKEHOLDERS</b></font><br/>{users}'),
        ]
        table = Table([cells], colWidths=[6.0 * cm, 14.0 * cm, content_width - 20.0 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5FAFE")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#A9C8DF")),
            ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C7DCEB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return [table, Spacer(1, 2)]

    def flow_box(title, items, color, pale):
        boxes = [p(f'<font color="#{color}"><b>{item}</b></font>', "tiny") for item in items]
        nested = Table([boxes], colWidths=[9.8 * cm / len(items)] * len(items))
        nested.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{pale}")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{color}")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return [p(f'<font color="#{color}"><b>{title.upper()}</b></font>', "tiny"), nested]

    expanded_rows = False

    def layer(number, name, question, color, pale, section_title, section_body, flow_title, flow_items, notes):
        left = p(f'<font color="#FFFFFF" size="9"><b>{number}</b></font><br/><font color="#FFFFFF"><b>{name.upper()}</b></font><br/><font color="#FFFFFF" size="5.5">{question}</font>', "white")
        objective = p(f'<font color="#{color}"><b>{section_title.upper()}</b></font><br/>{section_body}')
        process = flow_box(flow_title, flow_items, color, pale)
        note_cell = p(f'<font color="#{color}"><b>EXAMPLES / NOTES</b></font><br/>{notes}', "tiny")
        row_heights = [(3.0 if expanded_rows else 2.15) * cm]
        table = Table(
            [[left, objective, process, note_cell]],
            colWidths=[5.0 * cm, 8.3 * cm, 10.0 * cm, content_width - 23.3 * cm],
            rowHeights=row_heights,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(f"#{color}")),
            ("BACKGROUND", (1, 0), (1, 0), colors.white),
            ("BACKGROUND", (2, 0), (2, 0), colors.HexColor(f"#{pale}")),
            ("BACKGROUND", (3, 0), (3, 0), colors.HexColor("#F8FBFD")),
            ("BOX", (0, 0), (-1, -1), 0.55, colors.HexColor(f"#{color}")),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C6D8E5")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return [table, Spacer(1, 2)]

    def footer(principles, technologies, integrations):
        cells = [p(f'<font color="#0A2F5A"><b>ARCHITECTURE PRINCIPLES</b></font><br/>{principles}', "tiny"), p(f'<font color="#0A2F5A"><b>KEY TECHNOLOGIES</b></font><br/>{technologies}', "tiny"), p(f'<font color="#0A2F5A"><b>EXTERNAL INTEGRATIONS</b></font><br/>{integrations}', "tiny")]
        table = Table([cells], colWidths=[content_width / 3] * 3)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F7FA")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C7DCEB")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        return [table]

    def page_number(canvas, doc):
        canvas.saveState(); canvas.setFillColor(colors.HexColor("#60778C")); canvas.setFont("Helvetica", 6)
        canvas.drawRightString(page_width - 0.55 * cm, 0.18 * cm, f"PAGE {doc.page} OF 5")
        canvas.restoreState()

    summary = "AegisDesk clarifies incomplete requests, finds approved answers and verified solutions, recommends the right department and keeps people responsible for final decisions."
    users = "Requesters: every department<br/>Resolvers: IT, Data, Operations, Finance, HR, Security, Ecommerce<br/>Owners: service, knowledge, security and privacy"
    story = []
    story += header("ARCHITECTURE - LLM AGENTIC APPLICATION", "How the AegisDesk layers interact to deliver faster, safer cross-department service", "ARCHITECTURE MAP")
    story += overview(summary, users)
    story += layer("1", "Business Layer", "Why does this exist?", LAYER_COLORS["business"], LAYER_PALES["business"], "Objectives & value", "Complete requests<br/>Correct routing<br/>Faster answers<br/>Reusable knowledge", "Business process", ["REQUEST", "CLARIFY", "ROUTE", "RESOLVE"], "Less delay<br/>Measure time + routing")
    story += layer("2", "Agentic / Application Layer", "What do the agents do?", LAYER_COLORS["agentic"], LAYER_PALES["agentic"], "Agents & orchestration", "Five specialist roles<br/>Ordered handoffs<br/>Human approval", "Agent workflow", ["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY"], "Stop, ask or escalate<br/>No autonomous send")
    story += layer("3", "LLM Layer", "How is intelligence provided?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "LLM utilisation", "GPT-5.6 Terra<br/>Understands requests<br/>Asks missing questions<br/>Drafts grounded answers", "LLM components", ["REQUEST", "GPT-5.6", "PROMPT", "DRAFT"], "Language tasks only<br/>Human reviews output")
    story += layer("4", "Data Layer", "What data is used?", LAYER_COLORS["data"], LAYER_PALES["data"], "Data management", "SQLite ticket database<br/>Approved company documents<br/>Verified resolutions", "Data flow", ["SOURCES", "SEARCH", "EVIDENCE", "ANSWER"], "RAG, not training<br/>Permission-aware")
    story += layer("5", "Governance, Risk & Compliance", "How do we manage risk?", LAYER_COLORS["governance"], LAYER_PALES["governance"], "Governance & controls", "Privacy protection<br/>Evidence and confidence<br/>Human accountability", "Key controls", ["PRIVACY", "SECURITY", "QUALITY", "APPROVAL"], "Prevent unsafe answers<br/>Keep audit trail")
    story += layer("6", "Operations & Monitoring", "How do we run and improve it?", LAYER_COLORS["operations"], LAYER_PALES["operations"], "Operations & monitoring", "Business value<br/>LLM quality<br/>Service health<br/>Continuous improvement", "Improvement loop", ["COLLECT", "REVIEW", "IMPROVE", "RELEASE"], "Measure results<br/>Rollback when needed")
    story += footer("Scalable · Secure · Explainable · Human-centred", "Streamlit · GPT-5.6 Terra · SQLite · RAG", "Confluence · Company case system · SSO · Collaboration tools")
    story += [PageBreak()]

    expanded_rows = True
    story += header("BUSINESS & AGENTIC DESIGN", "How AegisDesk turns an unclear request into coordinated action", "PAGE 2 OF 5")
    story += overview("The whole company uses one platform. A department can request help in one workflow and resolve another request in the next.", users)
    story += layer("1", "Business Layer", "Why does this exist?", LAYER_COLORS["business"], LAYER_PALES["business"], "Objectives & value", "Reduce clarification loops<br/>Prevent wrong-team transfers<br/>Prioritise high-impact work<br/>Reuse company knowledge", "Process impacted", ["DESCRIBE", "COMPLETE", "ASSIGN", "ANSWER", "LEARN"], "KPIs:<br/>Completeness<br/>Routing accuracy<br/>Response + resolution time")
    story += layer("1A", "Business Value", "What changes for teams?", LAYER_COLORS["business"], LAYER_PALES["business"], "Value by user", "Requester: faster guidance<br/>Resolver: complete context<br/>Manager: visible ownership<br/>Company: reusable learning", "Value chain", ["LESS SEARCH", "FEWER TRANSFERS", "FASTER ACTION", "BETTER SERVICE"], "Investor value:<br/>Scales across departments<br/>Uses existing knowledge")
    story += layer("2", "Agentic / Application Layer", "What do the agents do?", LAYER_COLORS["agentic"], LAYER_PALES["agentic"], "Agent ecosystem", "Safety protects<br/>Triage understands<br/>Knowledge searches<br/>Resolution drafts<br/>Quality validates", "Example workflow", ["SAFETY", "TRIAGE", "KNOWLEDGE", "RESOLUTION", "QUALITY", "HUMAN"], "Orchestrator controls order<br/>Agents have separate roles")
    story += layer("2A", "Controlled Autonomy", "Who decides?", LAYER_COLORS["agentic"], LAYER_PALES["agentic"], "Decision rules", "Stop unsafe input<br/>Ask for missing facts<br/>Escalate weak evidence<br/>Route to approved teams", "Human-in-the-loop", ["AI RECOMMENDS", "PERSON REVIEWS", "EDIT / REASSIGN", "APPROVE / ESCALATE"], "No autonomous send<br/>No self-approval<br/>Visible trace")
    story += footer("Value first · Clear ownership · Controlled autonomy", "Employee Portal · Support Workspace · Five agents", "Company departments · Knowledge sources · Case systems")
    story += [PageBreak()]

    story += header("LLM LAYER", "Where GPT-5.6 Terra adds language intelligence - and where it does not", "PAGE 3 OF 5")
    story += overview("AegisDesk uses GPT-5.6 Terra through the OpenAI Responses API to read normal business language and prepare structured recommendations.", "Used by: Triage and Resolution Agents<br/>Reviewed by: Risk & Quality Agent and an authorised person")
    story += layer("3", "LLM Layer", "How is intelligence provided?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "LLM utilisation", "Understand the request<br/>Summarise the need<br/>Find missing information<br/>Propose classification<br/>Draft an answer", "LLM components", ["GPT-5.6 TERRA", "ROLE PROMPT", "COMPANY EVIDENCE", "RESULT"], "Language tasks only<br/>Evidence comes from RAG")
    story += layer("3A", "Triage Use", "How is the request understood?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "Input & result", "Input: safe request<br/>Output: summary, questions, category, priority and team recommendation", "Triage flow", ["REQUEST", "UNDERSTAND", "CLARIFY", "CLASSIFY", "ROUTE"], "Less manual reading<br/>Fewer clarification messages")
    story += layer("3B", "Resolution Use", "How is the answer prepared?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "Input & result", "Input: request + approved evidence<br/>Output: clear draft with sources and uncertainty", "Resolution flow", ["EVIDENCE", "DRAFT", "CHECK", "HUMAN REVIEW"], "No evidence = escalate<br/>The LLM must not invent")
    story += layer("3C", "LLM Boundary", "What can it not do?", LAYER_COLORS["llm"], LAYER_PALES["llm"], "Controlled responsibility", "Cannot send the final answer<br/>Cannot grant access<br/>Cannot close a ticket<br/>Cannot approve itself", "Control flow", ["LLM SUGGESTS", "RULES CHECK", "PERSON DECIDES"], "If OpenAI is unavailable:<br/>show failure<br/>do not fake live output")
    story += footer("Useful · Grounded · Reviewable · Human-controlled", "GPT-5.6 Terra · OpenAI Responses API", "Approved company evidence supplied through RAG")
    story += [PageBreak()]

    story += header("DATA LAYER", "Where every ticket, document and verified solution lives", "PAGE 4 OF 5")
    story += overview("The Data Layer gives AegisDesk memory. It stores operational tickets and makes approved company knowledge searchable.", "Owners: service teams and knowledge owners<br/>Access: authorised requesters and resolver teams")
    story += layer("4", "Data Layer", "What data is used?", LAYER_COLORS["data"], LAYER_PALES["data"], "Ticket database", "SQLite stores:<br/>Ticket description<br/>Requester and owner<br/>Status and priority<br/>SLA and activity<br/>Verified resolution", "Ticket life cycle", ["CREATE", "STORE", "ASSIGN", "UPDATE", "RESOLVE"], "Current database:<br/>data/aegisdesk.db")
    story += layer("4A", "Knowledge Sources", "Where do answers come from?", LAYER_COLORS["data"], LAYER_PALES["data"], "Approved knowledge", "Company procedures<br/>Operational guidance<br/>Data definitions<br/>Technical documentation<br/>Verified past solutions", "Source flow", ["CONFLUENCE / DOCS", "APPROVE", "INDEX", "SEARCH"], "Production sources keep owner, date and permissions")
    story += layer("4B", "RAG", "How is evidence found?", LAYER_COLORS["data"], LAYER_PALES["data"], "RAG in simple words", "Search first<br/>Select relevant evidence<br/>Give evidence to GPT-5.6 Terra<br/>Show the source", "RAG flow", ["REQUEST", "SEARCH", "BEST EVIDENCE", "LLM DRAFT", "CITATION"], "RAG does not retrain GPT<br/>It supplies current context")
    story += layer("4C", "Data Protection", "Who can see what?", LAYER_COLORS["data"], LAYER_PALES["data"], "Data controls", "User permissions<br/>Approved-source status<br/>Sensitive-data reduction<br/>Retention rules<br/>Human verification", "Access flow", ["IDENTIFY USER", "CHECK ACCESS", "FILTER SOURCES", "RETURN EVIDENCE"], "Future:<br/>SSO + managed database<br/>permission-aware Confluence")
    story += footer("Reliable · Permission-aware · Current · Reusable", "SQLite today · Managed database in production · RAG", "Confluence · Company documents · Existing tickets")
    story += [PageBreak()]

    story += header("GOVERNANCE, OPERATIONS & MONITORING", "How AegisDesk stays trustworthy and improves over time", "PAGE 5 OF 5")
    story += overview("AegisDesk combines safeguards, clear ownership and continuous measurement. AI recommendations remain visible and reversible.", "Owners: product, service, security, privacy, knowledge and platform operations")
    story += layer("5", "Governance, Risk & Compliance", "How do we manage risk?", LAYER_COLORS["governance"], LAYER_PALES["governance"], "Risks & controls", "Wrong answer: require evidence<br/>Sensitive data: protect<br/>Unsafe instruction: stop<br/>Wrong routing: reassign<br/>Over-trust: human review", "Key controls", ["PRIVACY", "SECURITY", "EVIDENCE", "QUALITY", "APPROVAL"], "Important actions have an owner and audit trail")
    story += layer("5A", "Human Accountability", "Who is responsible?", LAYER_COLORS["governance"], LAYER_PALES["governance"], "Ownership", "Knowledge owners approve sources<br/>Security/privacy approve controls<br/>Service owners measure outcomes<br/>People release answers", "Decision flow", ["AI PROPOSES", "CONTROL CHECK", "HUMAN DECIDES", "AUDIT"], "No certification claim without company assessment")
    story += layer("6", "Operations & Monitoring", "How do we run it?", LAYER_COLORS["operations"], LAYER_PALES["operations"], "What we monitor", "Business: time and routing<br/>LLM: usefulness and failures<br/>RAG: source relevance<br/>Service: speed and availability<br/>Safety: blocked events", "Observability loop", ["COLLECT", "MONITOR", "ALERT", "IMPROVE"], "Improve only after review and testing")
    story += layer("6A", "Continuous Improvement", "How does it improve?", LAYER_COLORS["operations"], LAYER_PALES["operations"], "Improvement cycle", "Review outcomes<br/>Find the cause<br/>Change one component<br/>Test quality and safety<br/>Release or rollback", "Controlled change", ["MEASURE", "LEARN", "TEST", "APPROVE", "RELEASE"], "Scale when value is demonstrated")
    story += footer("References: course assignment · OpenAI · Atlassian · NIST AI RMF · OWASP GenAI", "Streamlit · GPT-5.6 Terra · SQLite · scikit-learn", "Confluence · SSO · Company case system · Collaboration channels")
    story += [p("Disclosure: Generative AI assisted implementation and drafting. The student reviewed the work. Company deployment requires authorised data, credentials, permissions, security review and measured results.", "tiny")]
    pdf.build(story, onFirstPage=page_number, onLaterPages=page_number)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    pptx_path = OUTPUT_DIR / "AegisDesk_AI_One_Pager_v6.pptx"
    docx_path = OUTPUT_DIR / "AegisDesk_AI_Architecture_Template_Style_v8.docx"
    pdf_path = OUTPUT_DIR / "AegisDesk_AI_Architecture_Template_Style_v8.pdf"
    pptx_status = "Generated"
    try:
        build_one_pager(pptx_path)
    except PermissionError:
        pptx_status = "Preserved locked PPTX; close Microsoft PowerPoint to regenerate"
        print(f"Preserved locked PPTX (close Microsoft PowerPoint to regenerate): {pptx_path}")
    build_template_architecture_pdf(pdf_path)
    docx_status = "Generated"
    try:
        build_template_architecture_document(docx_path)
    except PermissionError:
        docx_status = "Skipped locked DOCX; close Microsoft Word to regenerate"
        print(f"Skipped locked DOCX (close Microsoft Word to regenerate): {docx_path}")
    print(f"{pptx_status}: {pptx_path}")
    print(f"{docx_status}: {docx_path}")
    print(f"Generated: {pdf_path}")


if __name__ == "__main__":
    main()
