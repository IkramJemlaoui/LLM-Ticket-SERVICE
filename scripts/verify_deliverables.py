from pathlib import Path

from docx import Document
from pptx import Presentation


root = Path(__file__).resolve().parents[1]
pptx_path = root / "deliverables" / "AegisDesk_AI_One_Pager_v6.pptx"
docx_path = root / "deliverables" / "AegisDesk_AI_Architecture_Template_Style_v8.docx"
pdf_path = root / "deliverables" / "AegisDesk_AI_Architecture_Template_Style_v8.pdf"

try:
    presentation = Presentation(pptx_path)
except PermissionError:
    presentation = None
document = Document(docx_path)
slide_text = "\n".join(
    shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text_frame")
) if presentation else ""
required_sections = [
    "APPLICATION OVERVIEW",
    "KEY CHALLENGES ADDRESSED",
    "CORE CAPABILITIES (SERVICES DELIVERED)",
    "DIFFERENTIATION",
]
section_positions = [slide_text.index(section) for section in required_sections] if presentation else []
document_text = "\n".join(
    [paragraph.text for paragraph in document.paragraphs]
    + [paragraph.text for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
)
page_breaks = sum(
    1
    for element in document.element.body.iter()
    if element.tag.endswith("}br")
    and element.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page"
)

if presentation:
    assert len(presentation.slides) == 1, "The one-pager must contain exactly one slide."
    assert section_positions == sorted(section_positions), "Required one-pager sections must appear in order."
    assert "General use cases for testing" not in slide_text
assert page_breaks == 4, "The architecture source must contain four explicit breaks for five pages."
for layer in ["Business layer", "Agentic / Application layer", "LLM layer", "Data layer", "Governance, Risk", "Operations & Monitoring"]:
    assert layer.upper() in document_text.upper(), f"Missing architecture layer label: {layer}"
if presentation:
    assert presentation.core_properties.author == "Ikram Jemlaoui"
assert document.core_properties.author == "Ikram Jemlaoui"
pdf_bytes = pdf_path.read_bytes()
pdf_page_count = pdf_bytes.count(b"/Type /Page\n")
assert pdf_page_count == 5, f"Architecture PDF must be exactly five pages, found {pdf_page_count}."

if presentation:
    print(f"PPTX: 1 slide, {len(presentation.slides[0].shapes)} shapes, {pptx_path.stat().st_size} bytes")
else:
    print("PPTX: locked by PowerPoint; preserved and skipped during this verification run")
print(f"DOCX: 5 page sections, {len(document.tables)} tables, {docx_path.stat().st_size} bytes")
print(f"PDF: {pdf_page_count} pages, {pdf_path.stat().st_size} bytes")
