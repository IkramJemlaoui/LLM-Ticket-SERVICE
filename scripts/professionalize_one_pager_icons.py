"""Replace v7 emoji artwork with restrained native PowerPoint vector icons."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "AegisDesk_AI_One_Pager_v7.pptx"
OUTPUT = ROOT / "deliverables" / "AegisDesk_AI_One_Pager_v9_Ollama.pptx"

BLUE = RGBColor(37, 91, 173)
VIOLET = RGBColor(113, 55, 166)
TEAL = RGBColor(7, 143, 145)
ORANGE = RGBColor(231, 138, 18)
PURPLE = RGBColor(102, 69, 170)
WHITE = RGBColor(255, 255, 255)


ICON_MAP = {
    "💡": "idea",
    "🏷": "tag",
    "👥": "users",
    "🎯": "target",
    "💎": "diamond",
    "⚠": "warning",
    "🔀": "route",
    "📚": "knowledge",
    "🛡": "shield",
    "📝": "document",
    "🧭": "compass",
    "🔎": "search",
    "✅": "check",
    "✨": "star",
    "🤖": "network",
    "🔌": "integration",
    "📈": "chart",
    "✍": "pencil",
    "📊": "kpi",
    "💰": "roi",
    "⏱": "clock",
    "🧩": "foundation",
}


def accent_for(x: float, y: float) -> RGBColor:
    if y < 1.15:
        return PURPLE
    if y > 6.3:
        return TEAL
    if y < 3.75:
        if x < 4.45:
            return BLUE
        if x < 8.78:
            return VIOLET
        return TEAL
    if x < 4.45:
        return ORANGE
    if x < 8.78:
        return BLUE
    return PURPLE


def outline(shape, color, width=1.15, fill=None):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def auto(slide, kind, x, y, w, h, color, fill=None, rotation=0):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    outline(shape, color, fill=fill)
    shape.rotation = rotation
    return shape


def line(slide, x1, y1, x2, y2, color, width=1.15):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def label(slide, x, y, w, h, text, color, size=6.5):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.clear()
    shape.text_frame.margin_left = shape.text_frame.margin_right = 0
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos Display"
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return shape


def draw_icon(slide, x, y, w, h, kind, color):
    # Use a square drawing area centred inside the original icon placeholder.
    size = min(w, h) * 0.88
    x = x + (w - size) / 2
    y = y + (h - size) / 2
    w = h = size
    cx, cy = x + w / 2, y + h / 2

    if kind == "idea":
        auto(slide, MSO_SHAPE.OVAL, x + 0.22*w, y, 0.56*w, 0.58*h, color)
        line(slide, x + 0.34*w, y + 0.62*h, x + 0.66*w, y + 0.62*h, color)
        line(slide, x + 0.38*w, y + 0.76*h, x + 0.62*w, y + 0.76*h, color)
        line(slide, cx, y + 0.58*h, cx, y + 0.69*h, color)
    elif kind == "tag":
        auto(slide, MSO_SHAPE.PENTAGON, x + 0.08*w, y + 0.12*h, 0.82*w, 0.72*h, color, rotation=90)
        auto(slide, MSO_SHAPE.OVAL, x + 0.62*w, y + 0.34*h, 0.12*w, 0.12*h, color)
    elif kind == "users":
        auto(slide, MSO_SHAPE.OVAL, x + 0.12*w, y + 0.10*h, 0.28*w, 0.28*h, color)
        auto(slide, MSO_SHAPE.OVAL, x + 0.58*w, y + 0.10*h, 0.28*w, 0.28*h, color)
        auto(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.02*w, y + 0.48*h, 0.46*w, 0.35*h, color)
        auto(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.52*w, y + 0.48*h, 0.46*w, 0.35*h, color)
    elif kind == "target":
        auto(slide, MSO_SHAPE.OVAL, x + 0.04*w, y + 0.04*h, 0.92*w, 0.92*h, color)
        auto(slide, MSO_SHAPE.OVAL, x + 0.27*w, y + 0.27*h, 0.46*w, 0.46*h, color)
        auto(slide, MSO_SHAPE.OVAL, x + 0.44*w, y + 0.44*h, 0.12*w, 0.12*h, color, fill=color)
    elif kind == "diamond":
        auto(slide, MSO_SHAPE.DIAMOND, x + 0.07*w, y + 0.12*h, 0.86*w, 0.72*h, color)
        line(slide, x + 0.28*w, y + 0.14*h, cx, y + 0.84*h, color, 0.8)
        line(slide, x + 0.72*w, y + 0.14*h, cx, y + 0.84*h, color, 0.8)
    elif kind == "warning":
        auto(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x + 0.05*w, y + 0.03*h, 0.9*w, 0.88*h, color)
        line(slide, cx, y + 0.34*h, cx, y + 0.60*h, color, 1.4)
        auto(slide, MSO_SHAPE.OVAL, x + 0.46*w, y + 0.69*h, 0.08*w, 0.08*h, color, fill=color)
    elif kind == "route":
        arrow1 = auto(slide, MSO_SHAPE.RIGHT_ARROW, x + 0.05*w, y + 0.10*h, 0.88*w, 0.30*h, color, fill=color)
        arrow1.line.fill.background()
        arrow2 = auto(slide, MSO_SHAPE.LEFT_ARROW, x + 0.07*w, y + 0.58*h, 0.88*w, 0.30*h, color, fill=color)
        arrow2.line.fill.background()
    elif kind == "knowledge":
        auto(slide, MSO_SHAPE.RECTANGLE, x + 0.07*w, y + 0.18*h, 0.36*w, 0.66*h, color)
        auto(slide, MSO_SHAPE.RECTANGLE, x + 0.55*w, y + 0.18*h, 0.36*w, 0.66*h, color)
        line(slide, cx, y + 0.25*h, cx, y + 0.82*h, color)
        line(slide, x + 0.18*w, y + 0.36*h, x + 0.36*w, y + 0.36*h, color, 0.8)
        line(slide, x + 0.64*w, y + 0.36*h, x + 0.82*w, y + 0.36*h, color, 0.8)
    elif kind == "shield":
        auto(slide, MSO_SHAPE.PENTAGON, x + 0.14*w, y + 0.05*h, 0.72*w, 0.84*h, color, rotation=180)
        line(slide, x + 0.32*w, y + 0.46*h, x + 0.46*w, y + 0.62*h, color)
        line(slide, x + 0.46*w, y + 0.62*h, x + 0.70*w, y + 0.31*h, color)
    elif kind == "document":
        auto(slide, MSO_SHAPE.FOLDED_CORNER, x + 0.13*w, y + 0.03*h, 0.72*w, 0.9*h, color)
        for offset in (0.38, 0.54, 0.70):
            line(slide, x + 0.27*w, y + offset*h, x + 0.68*w, y + offset*h, color, 0.8)
    elif kind == "compass":
        auto(slide, MSO_SHAPE.OVAL, x + 0.04*w, y + 0.04*h, 0.92*w, 0.92*h, color)
        auto(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x + 0.35*w, y + 0.12*h, 0.30*w, 0.72*h, color, fill=color, rotation=25)
        auto(slide, MSO_SHAPE.OVAL, x + 0.43*w, y + 0.43*h, 0.14*w, 0.14*h, WHITE, fill=WHITE)
    elif kind == "search":
        auto(slide, MSO_SHAPE.OVAL, x + 0.08*w, y + 0.06*h, 0.58*w, 0.58*h, color)
        line(slide, x + 0.60*w, y + 0.60*h, x + 0.91*w, y + 0.91*h, color, 1.5)
    elif kind == "check":
        auto(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.05*w, y + 0.05*h, 0.9*w, 0.9*h, color)
        line(slide, x + 0.22*w, y + 0.51*h, x + 0.43*w, y + 0.72*h, color, 1.6)
        line(slide, x + 0.43*w, y + 0.72*h, x + 0.80*w, y + 0.29*h, color, 1.6)
    elif kind == "star":
        auto(slide, MSO_SHAPE.STAR_5_POINT, x + 0.05*w, y + 0.05*h, 0.9*w, 0.9*h, color)
    elif kind == "network":
        points = [(cx, y + 0.14*h), (x + 0.18*w, y + 0.75*h), (x + 0.82*w, y + 0.75*h)]
        line(slide, *points[0], *points[1], color)
        line(slide, *points[0], *points[2], color)
        line(slide, *points[1], *points[2], color)
        for px, py in points:
            auto(slide, MSO_SHAPE.OVAL, px - 0.11*w, py - 0.11*h, 0.22*w, 0.22*h, color, fill=WHITE)
    elif kind == "integration":
        auto(slide, MSO_SHAPE.OVAL, x + 0.05*w, y + 0.25*h, 0.58*w, 0.34*h, color, rotation=-35)
        auto(slide, MSO_SHAPE.OVAL, x + 0.38*w, y + 0.40*h, 0.58*w, 0.34*h, color, rotation=-35)
    elif kind in {"chart", "kpi"}:
        for bx, bh in ((0.10, 0.32), (0.38, 0.52), (0.66, 0.76)):
            bar = auto(slide, MSO_SHAPE.RECTANGLE, x + bx*w, y + (0.9-bh)*h, 0.18*w, bh*h, color, fill=color)
            bar.line.fill.background()
        line(slide, x + 0.06*w, y + 0.92*h, x + 0.92*w, y + 0.92*h, color)
    elif kind == "pencil":
        body = auto(slide, MSO_SHAPE.RECTANGLE, x + 0.18*w, y + 0.38*h, 0.68*w, 0.20*h, color, fill=color, rotation=-42)
        body.line.fill.background()
        auto(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x + 0.04*w, y + 0.67*h, 0.25*w, 0.22*h, color, fill=color, rotation=45)
    elif kind == "roi":
        for ox, oy in ((0.10, 0.43), (0.38, 0.18), (0.66, 0.43)):
            auto(slide, MSO_SHAPE.OVAL, x + ox*w, y + oy*h, 0.28*w, 0.28*h, color)
        label(slide, x + 0.29*w, y + 0.43*h, 0.42*w, 0.26*h, "$", color, max(5, 7*w/0.3))
    elif kind == "clock":
        auto(slide, MSO_SHAPE.OVAL, x + 0.04*w, y + 0.04*h, 0.92*w, 0.92*h, color)
        line(slide, cx, cy, cx, y + 0.23*h, color, 1.3)
        line(slide, cx, cy, x + 0.73*w, y + 0.61*h, color, 1.3)
        auto(slide, MSO_SHAPE.OVAL, x + 0.45*w, y + 0.45*h, 0.10*w, 0.10*h, color, fill=color)
    elif kind == "foundation":
        for ox, oy in ((0.08, 0.20), (0.40, 0.05), (0.40, 0.47), (0.72, 0.20)):
            auto(slide, MSO_SHAPE.HEXAGON, x + ox*w, y + oy*h, 0.27*w, 0.27*h, color)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    prs = Presentation(SOURCE)
    if len(prs.slides) != 1:
        raise RuntimeError("Expected exactly one slide.")
    slide = prs.slides[0]
    replacements = []

    for shape in list(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if "GPT-5.6 Terra" in run.text:
                    run.text = run.text.replace(
                        "GPT-5.6 Terra · OpenAI Responses API",
                        "Qwen 2.5 · local Ollama inference",
                    )
        original = shape.text.strip()
        normalized = original.replace("\ufe0f", "")
        if normalized.startswith("👤 HUMAN APPROVAL"):
            shape.text_frame.paragraphs[0].text = "HUMAN APPROVAL · FINAL AUTHORITY"
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(5.3)
            paragraph.font.bold = True
            paragraph.font.color.rgb = PURPLE
            continue
        kind = ICON_MAP.get(normalized)
        if kind is None:
            continue
        replacements.append((shape, kind))
        shape.text_frame.clear()

    for shape, kind in replacements:
        x, y = shape.left / 914400, shape.top / 914400
        w, h = shape.width / 914400, shape.height / 914400
        draw_icon(slide, x, y, w, h, kind, accent_for(x, y))

    prs.core_properties.title = "AegisDesk AI — Six-Section One-Pager with Local Ollama LLM"
    prs.core_properties.comments = "All emoji artwork replaced with native PowerPoint vector icons."
    prs.save(OUTPUT)
    print(f"Generated: {OUTPUT}")
    print(f"Replaced emoji elements: {len(replacements)}")
    print(f"Slides: {len(prs.slides)}; shapes: {len(slide.shapes)}")


if __name__ == "__main__":
    main()
